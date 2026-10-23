"""
SoftProdigy Snow Agent — endpoints the UI depends on that were missing from v5:

  POST /snow/extract/universal          any input -> user stories (text, PDF, DOCX, XLSX/CSV, audio, image)
  POST /snow/stories/approve-and-deploy approved stories -> incident / problem / change_request / rm_story records
  POST /snow/health-scan                live health scan in the UI's format (sn_health)
  POST /snow/universal/chat             general ServiceNow assistant chat
  POST /snow/jira/push                  stories -> Jira issues (parity with the SoftProdigy platform's pushToJira)
"""

from __future__ import annotations

import asyncio
import csv
import io
import json
import logging
import os
import re
from collections import Counter
from typing import Any, Dict, List, Optional, Tuple

import requests
from fastapi import APIRouter, File, Form, HTTPException, UploadFile

import docx_render as dr
import groq_client
import sn_client as snc
import requirement_pipeline as rp
import sn_health
import story_build
from sn_client import SNError

logger = logging.getLogger(__name__)
snow_universal_router = APIRouter(tags=["Snow Universal"])

AUDIO = {"mp3", "mp4", "wav", "m4a", "ogg", "flac", "webm", "mpeg", "mpga"}
IMAGE = {"png", "jpg", "jpeg", "webp", "bmp", "gif"}
MAX_AUDIO_MB = 25  # Groq Whisper upload limit
CHUNK = 6000                                                   # characters per AI call (split on line breaks)
MAX_INPUT = int(os.getenv("AGENT_MAX_INPUT_CHARS", "300000"))  # ~100 pages of notes


# ════════════════════════════════════════════════════════════════════════════
# File -> text
# ════════════════════════════════════════════════════════════════════════════

def _pdf_text(b: bytes) -> str:
    try:
        import PyPDF2
        text = "\n".join(p.extract_text() or "" for p in PyPDF2.PdfReader(io.BytesIO(b)).pages)
        if text.strip():
            return text
    except Exception:
        pass
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(b)) as pdf:
            return "\n".join(p.extract_text() or "" for p in pdf.pages)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Could not read the PDF (scanned PDFs need OCR): {e}")


def _docx_text(b: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(b))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:  # v5 ignored tables — requirements often live in them
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _xlsx_text(b: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h or f"col{i}") for i, h in enumerate(rows[0])]
        out.append(f"## Sheet: {ws.title}")
        for r in rows[1:]:
            vals = {header[i]: v for i, v in enumerate(r) if i < len(header) and v not in (None, "")}
            if vals:
                out.append("; ".join(f"{k}: {v}" for k, v in vals.items()))
    return "\n".join(out)


def _csv_text(b: bytes) -> str:
    text = b.decode("utf-8-sig", errors="replace")
    rows = list(csv.DictReader(io.StringIO(text)))
    return "\n".join("; ".join(f"{k}: {v}" for k, v in r.items() if v) for r in rows)


# ── Structured story sheets (e.g. User Story ID | Summary | Description | Acceptance Criteria) ──
_COLS = {
    "id": ("user story id", "story id", "id", "key", "issue key", "us id"),
    "title": ("summary", "title", "story title", "name", "feature"),
    "description": ("description", "user story", "story", "requirement"),
    "acceptance_criteria": ("acceptance criteria", "acceptance", "criteria", "ac"),
    "priority": ("priority",),
    "module": ("module", "product", "component"),
    "story_points": ("story points", "points", "estimate"),
}
_AS_A = re.compile(r"^\s*as an?\s+(?P<as_a>.+?),\s*i want\s+(?:to\s+)?(?P<i_want>.+?)(?:,?\s+so that\s+(?P<so_that>.+?))?\.?\s*$",
                   re.I | re.S)


def _map_header(header: List[str]) -> Dict[str, int]:
    norm = [re.sub(r"[^a-z ]", "", str(h or "").lower()).strip() for h in header]
    found: Dict[str, int] = {}
    for key, names in _COLS.items():
        for i, h in enumerate(norm):
            if i not in found.values() and (h in names or any(h.startswith(n) for n in names if len(n) > 3)):
                found[key] = i
                break
    return found


def _rows_from_bytes(b: bytes, ext: str) -> List[List[Any]]:
    if ext == "csv":
        return list(csv.reader(io.StringIO(b.decode("utf-8-sig", errors="replace"))))
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True).worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


_MODULE_WORDS = [("ITOM", ("discovery", "mid server", "event management", "service mapping", "alert")),
                 ("CMDB", ("cmdb", "csdm", " ci ", "configuration item", "reconcil")),
                 ("Catalog", ("catalog", "request item", "portal")), ("HRSD", ("hr ", "hrms", "onboard")),
                 ("SecOps", ("vulnerab", "security incident")), ("GRC", ("compliance", "audit", "risk")),
                 ("Integration", ("integrat", "api", "jira", "email", "microsoft 365")),
                 ("ITSM", ("incident", "problem", "change", "sla", "assignment"))]


def _guess_module(text: str) -> str:
    t = f" {text.lower()} "
    for mod, words in _MODULE_WORDS:
        if any(w in t for w in words):
            return mod
    return "Platform"


def structured_stories(b: bytes, ext: str) -> Optional[Dict[str, Any]]:
    """If the sheet already has story columns, map rows 1:1 — no AI, nothing rewritten, IDs kept."""
    try:
        rows = [r for r in _rows_from_bytes(b, ext) if any(str(c or "").strip() for c in r)]
    except Exception:
        return None
    if len(rows) < 2:
        return None
    cols = _map_header(rows[0])
    if "acceptance_criteria" not in cols or not ({"title", "description"} & set(cols)):
        return None
    get = lambda r, k: str(r[cols[k]]).strip() if k in cols and cols[k] < len(r) and r[cols[k]] is not None else ""
    stories = []
    for n, r in enumerate(rows[1:], 1):
        title, desc = get(r, "title"), get(r, "description")
        if not (title or desc):
            continue
        m = _AS_A.match(desc) if desc else None
        pri = re.sub(r"[^a-z]", "", get(r, "priority").lower())
        pri = {"1": "critical", "2": "high", "3": "medium", "4": "low", "p1": "critical", "p2": "high"}.get(pri, pri)
        try:
            pts = int(float(get(r, "story_points") or 0)) or None
        except ValueError:
            pts = None
        stories.append({
            "id": get(r, "id") or f"US-{n:03d}", "title": title or desc[:80], "description": desc,
            "as_a": (m.group("as_a") if m else "") or "user", "i_want": (m.group("i_want") if m else "") or title or desc,
            "so_that": (m.group("so_that") if m else "") or "",
            "acceptance_criteria": dr.as_list(get(r, "acceptance_criteria")), "criteria_inferred": False,
            "priority": pri if pri in ("critical", "high", "medium", "low") else "medium",
            "module": get(r, "module") or _guess_module(f"{title} {desc}"), "story_type": "story",
            "story_points": pts or 3, "source_row": n + 1})
    if not stories:
        return None
    ids = [s["id"] for s in stories]
    prefix = re.sub(r"[_\-\s]*\d+$", "", ids[0]) if all(re.sub(r"[_\-\s]*\d+$", "", i) == re.sub(r"[_\-\s]*\d+$", "", ids[0]) for i in ids) else ""
    return {"project_title": prefix.replace("_", " ").strip() if prefix and prefix != "US" else "",
            "project_overview": "", "user_stories": stories, "total_stories": len(stories), "truncated": False,
            "mapping": {k: rows[0][v] for k, v in cols.items()}}


async def file_to_text(upload: UploadFile) -> Dict[str, Any]:
    data = await upload.read()
    name = (upload.filename or "file").lower()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""
    if not data:
        raise HTTPException(status_code=422, detail="The uploaded file is empty")
    if ext == "pdf":
        return {"type": "pdf", "text": await asyncio.to_thread(_pdf_text, data)}
    if ext in ("docx",):
        return {"type": "docx", "text": await asyncio.to_thread(_docx_text, data)}
    if ext == "doc":
        raise HTTPException(status_code=415, detail="Old .doc files are not supported — save as .docx or PDF")
    if ext in ("xlsx", "xlsm"):
        return {"type": "spreadsheet", "text": await asyncio.to_thread(_xlsx_text, data)}
    if ext == "xls":
        raise HTTPException(status_code=415, detail="Old .xls files are not supported — save as .xlsx or .csv")
    if ext == "csv":
        return {"type": "csv", "text": _csv_text(data)}
    if ext in AUDIO:
        if len(data) > MAX_AUDIO_MB * 1024 * 1024:
            raise HTTPException(status_code=413, detail=f"Audio is larger than {MAX_AUDIO_MB} MB (Whisper limit) — "
                                                        "split the recording or export it as a lower-bitrate mp3")
        try:
            text = await asyncio.to_thread(groq_client.transcribe, data, upload.filename or f"audio.{ext}")
        except groq_client.GroqError as e:
            raise HTTPException(status_code=503, detail=f"Transcription failed: {e}")
        return {"type": "audio", "text": text, "transcript": True}
    if ext in IMAGE:
        import base64
        mime = upload.content_type or f"image/{'jpeg' if ext == 'jpg' else ext}"
        try:
            text = await asyncio.to_thread(
                groq_client.describe_image, base64.b64encode(data).decode(), mime,
                "Transcribe ALL text on this whiteboard / screenshot / page exactly, then describe any diagram "
                "(boxes, arrows, swimlanes) as a numbered list of steps. Do not invent content.")
        except groq_client.GroqError as e:
            raise HTTPException(status_code=503, detail=f"Image reading failed: {e}")
        return {"type": "image", "text": text}
    if ext in ("txt", "md", "json", "eml", "html", "htm", "xml", "log", "") or \
            (upload.content_type or "").startswith("text/"):
        return {"type": "text", "text": data.decode("utf-8", errors="replace")}
    raise HTTPException(status_code=415, detail=f"Unsupported file type '.{ext}'")


# ════════════════════════════════════════════════════════════════════════════
# Text -> stories
# ════════════════════════════════════════════════════════════════════════════

_SYSTEM = ("You are a senior ServiceNow business analyst. You extract user stories ONLY from what the input says; "
           "you never invent requirements that are not in the input. Respond with JSON only.")

_PROMPT = """Extract every user story / requirement from the input below (it may be meeting notes, a transcript,
an email, a spreadsheet dump or a formal document). Merge duplicates. Keep the user's wording where possible.

For each story:
- title: short (max 80 chars)
- as_a / i_want / so_that
- acceptance_criteria: list of testable criteria (Given/When/Then when possible). Use the input's criteria;
  only if none are given, write the minimum criteria implied by the story and set "criteria_inferred": true.
- priority: critical | high | medium | low
- module: ITSM | ITOM | CMDB | HRSD | CSM | SecOps | GRC | ITAM | Catalog | FSM | Platform | Integration
- story_type: incident | problem | change | story   (incident/problem/change ONLY when the item is an actual
  operational issue or change to log, not a feature to build)
- story_points: 1,2,3,5,8,13

INPUT ({kind}):
\"\"\"
{text}
\"\"\"

There is NO limit on the number of stories: return every distinct story in this input (it may be 1 or 50).
Do not merge different requirements into one story and do not summarise several stories away.

Return JSON: {{"project_title": "...", "project_overview": "2-3 sentences", "user_stories": [{{...}}],
"complete": true}}   <- always end with "complete": true"""


def _norm_story(s: dict) -> Optional[dict]:
    if not isinstance(s, dict):
        return None
    title = dr.s(s.get("title") or s.get("i_want") or s.get("story")).strip()
    if not title:
        return None
    pri = re.sub(r"[^a-z]", "", str(s.get("priority", "medium")).lower())
    stype = str(s.get("story_type", "story")).lower()
    try:
        pts = int(float(s.get("story_points") or 3))
    except (TypeError, ValueError):
        pts = 3
    return {"title": title[:160], "description": dr.s(s.get("description")),
            "as_a": dr.s(s.get("as_a")) or "user", "i_want": dr.s(s.get("i_want")) or title,
            "so_that": dr.s(s.get("so_that")), "acceptance_criteria": dr.as_list(s.get("acceptance_criteria")),
            "criteria_inferred": bool(s.get("criteria_inferred")),
            "priority": pri if pri in ("critical", "high", "medium", "low") else "medium",
            "module": dr.s(s.get("module")) or "Platform",
            "story_type": stype if stype in ("incident", "problem", "change", "story") else "story",
            "story_points": pts}


async def text_to_stories(text: str, kind: str) -> Dict[str, Any]:
    text = (text or "").strip()
    if len(text) < 10:
        raise HTTPException(status_code=422, detail="Not enough text to extract stories from")
    def split_lines(t: str, size: int) -> List[str]:
        out, cur = [], ""
        for line in t.splitlines(keepends=True):   # split on line boundaries, never mid-record
            if cur and len(cur) + len(line) > size:
                out.append(cur)
                cur = ""
            cur += line
        if cur.strip():
            out.append(cur)
        return out

    chunks = split_lines(text[:MAX_INPUT], CHUNK)
    retries = {"split": 0}

    async def one(chunk: str, depth: int = 0) -> List[dict]:
        """One AI call. If the answer is cut off (no "complete": true / unreadable), split the text and retry."""
        msgs = [{"role": "system", "content": _SYSTEM + "\nRespond with a single valid JSON object only."},
                {"role": "user", "content": _PROMPT.format(kind=kind, text=chunk)}]
        try:
            # no automatic "repair": a cut-off answer must be detected, not patched (patching drops stories)
            raw = await asyncio.to_thread(groq_client.chat, msgs, 8000, 0.2, None, True)
        except groq_client.GroqError as e:
            raise HTTPException(status_code=503, detail=f"AI service unavailable: {e}")
        try:
            p = groq_client.parse_json(raw)
            p = p if isinstance(p, dict) else {"user_stories": p}
            complete = p.get("complete") is True
        except ValueError:
            p, complete = {}, False
        if not complete and depth < 3 and len(chunk) > 800:
            halves = split_lines(chunk, max(400, len(chunk) // 2 + 1))
            if len(halves) > 1:
                retries["split"] += 1
                parts = await asyncio.gather(*(one(h, depth + 1) for h in halves))
                return [x for part in parts for x in part]
        if not p:
            raise HTTPException(status_code=502, detail="AI returned an unreadable answer, please retry")
        return [{"_meta": {"project_title": p.get("project_title"), "project_overview": p.get("project_overview")}}] + \
            list(p.get("user_stories") or p.get("stories") or [])

    results = await asyncio.gather(*(one(c) for c in chunks))
    stories, seen = [], set()
    title = overview = ""
    for items in results:
        for s in items:
            if isinstance(s, dict) and "_meta" in s:
                title = title or dr.s(s["_meta"].get("project_title"))
                overview = overview or dr.s(s["_meta"].get("project_overview"))
                continue
            n = _norm_story(s)
            if not n:
                continue
            # only drop exact repeats (same title AND same need) — e.g. a story repeated across two chunks
            key = (re.sub(r"\W+", "", n["title"].lower()), re.sub(r"\W+", "", n["i_want"].lower()))
            if key not in seen:
                seen.add(key)
                stories.append(n)
    for i, s in enumerate(stories, 1):
        s["id"] = f"US-{i:03d}"
    return {"project_title": title, "project_overview": overview, "user_stories": stories,
            "total_stories": len(stories), "truncated": len(text) > MAX_INPUT,
            "ai_calls": {"chunks": len(chunks), "re_split_after_cut_off": retries["split"]}}


@snow_universal_router.post("/snow/extract/universal")
async def extract_universal(file: Optional[UploadFile] = File(None), text: Optional[str] = Form(None)):
    """Any input -> user stories. Send either `file` or `text` (multipart form)."""
    if file is not None and file.filename:
        ext = (file.filename.rsplit(".", 1)[-1] if "." in file.filename else "").lower()
        if ext in ("xlsx", "xlsm", "csv"):
            data = await file.read()
            direct = await asyncio.to_thread(structured_stories, data, "csv" if ext == "csv" else "xlsx")
            if direct:
                direct["source"] = {"type": "spreadsheet", "mode": "direct (no AI rewriting)",
                                    "filename": file.filename, "columns": direct.pop("mapping")}
                return direct
            await file.seek(0)
        src = await file_to_text(file)
    elif text and text.strip():
        src = {"type": "text", "text": text}
    else:
        raise HTTPException(status_code=400, detail="Send a file or some text")
    if not src["text"].strip():
        raise HTTPException(status_code=422, detail=f"No text could be read from the {src['type']}")
    out = await text_to_stories(src["text"], src["type"])
    out["source"] = {"type": src["type"], "characters": len(src["text"]),
                     "filename": getattr(file, "filename", None) if file else None}
    if src.get("transcript"):
        out["transcript"] = src["text"][:20000]
    out["_raw_text"] = src["text"][:20000]
    if not out["user_stories"]:
        raise HTTPException(status_code=422, detail="The AI found no user stories in this input")
    return out


# ════════════════════════════════════════════════════════════════════════════
# Stories -> ServiceNow records
# ════════════════════════════════════════════════════════════════════════════

# priority -> (impact, urgency); ServiceNow calculates priority from these on incident/problem
_IU = {"critical": ("1", "1"), "high": ("1", "2"), "medium": ("2", "2"), "low": ("3", "3")}
def _find_group(sn: snc.SNClient, name: str, cache: dict) -> Tuple[Optional[dict], str]:
    """Exact lookup of an active group by name (or sys_id). Returns (group, problem). No prefix or module guessing:
    a name that matches no group or several groups is a question for the user."""
    if name in cache:
        return cache[name]
    if snc.is_sys_id(name):
        rows = sn.query("sys_user_group", f"active=true^sys_id={name}", "sys_id,name", limit=2)
    else:
        rows = sn.query("sys_user_group", f"active=true^name={name}", "sys_id,name", limit=2)
    if len(rows) == 1:
        out = (rows[0], "")
    elif rows:
        out = (None, f"Assignment group '{name}' matches several active groups — give its sys_id")
    else:
        out = (None, f"No active assignment group named '{name}' exists on this instance — which group should be used?")
    cache[name] = out
    return out


@snow_universal_router.post("/snow/stories/approve-and-deploy")
async def approve_and_deploy(body: dict):
    """
    Create one ServiceNow record per approved story.
      story_type incident -> incident | problem -> problem | change -> change_request |
      story/task -> rm_story (Agile plugin) else task
    Stories that describe platform code (script include, business rule, client script, UI policy) also get that
    code built in the global scope (see story_build); set body.build_artifacts=false for tracking records only.
    With body.pipeline=true such stories are instead planned by requirement_pipeline (read-only) and returned as
    `requirement_plan`; the build happens only after the user confirms it via POST /snow/requirements/deploy.
    Requires body.approved == true (the UI sets it only after the user confirms each story).
    Assignment group: story.assignment_group, else body.group_mapping[module], else body.assignment_group — matched
    exactly; no group is guessed. An unresolvable or ambiguous group returns an error with needs_input=true.
    Response items carry `index` = position in the request, so the UI never mismatches results.
    """
    stories = body.get("user_stories") or []
    if not stories:
        raise HTTPException(status_code=400, detail="No stories provided")
    if body.get("approved") is not True:
        raise HTTPException(status_code=428, detail="Deployment not approved — send approved=true after the user "
                                                    "confirms (this call creates records in ServiceNow)")
    strict_planner = body.get("require_preflight") is True and body.get("pipeline") is True
    planner_stories = []
    if strict_planner:
        for index, story in enumerate(stories):
            normalized = _norm_story(story) or {}
            if story_build.is_buildable(normalized) or rp.looks_like_configuration(story_build.story_text(normalized)):
                planner_stories.append((index, story))
    strict_planner = bool(planner_stories)
    if strict_planner:
        invalid = []
        for index, story in planner_stories:
            criteria = story.get("acceptance_criteria") or story.get("criteria") or []
            if isinstance(criteria, str):
                criteria = criteria.splitlines()
            if not isinstance(criteria, list) or not any(str(item).strip() for item in criteria):
                invalid.append(index)
        if invalid:
            raise HTTPException(status_code=422, detail={"message": "Every Planner story needs acceptance criteria.",
                                                         "story_indexes": invalid})
    sn = snc.client_from(body)

    def work():
        try:
            sn.user()
        except SNError as e:
            raise e.to_http("ServiceNow login")
        readiness = None
        if strict_planner:
            from snow_endpoints import _planner_preflight_report
            readiness = _planner_preflight_report(sn, body)
            if not readiness["execution_ready"]:
                raise HTTPException(status_code=412, detail="Planner preflight blocked story processing: "
                                    + "; ".join(readiness["blockers"]))
        has_story = sn.table_exists("rm_story")
        build_artifacts = body.get("build_artifacts", True) is not False
        use_pipeline = body.get("pipeline") is True and build_artifacts
        groups: dict = {}
        created, errors = [], []
        for idx, raw in enumerate(stories):
            s = _norm_story(raw) or {}
            if not s:
                errors.append({"index": idx, "error": "story has no title"})
                continue
            stype = str(raw.get("story_type") or s["story_type"]).lower()
            retyped = ""
            if stype in ("incident", "problem", "change") and story_build.is_buildable(s):
                # A story that asks for a script include / client script / business rule / UI policy is platform
                # configuration — never an incident, problem or change, whatever words it uses
                retyped, stype = stype, "story"
            table = {"incident": "incident", "problem": "problem", "change": "change_request"}.get(
                stype, "rm_story" if has_story else "task")
            ac = s["acceptance_criteria"]
            desc = s["description"] or (f"As a {s['as_a']}, I want {s['i_want']}" +
                                        (f", so that {s['so_that']}." if s["so_that"] else "."))
            if raw.get("id"):
                desc = f"[{raw['id']}] " + desc
            if ac and table != "rm_story":
                desc += "\n\nAcceptance Criteria:\n" + "\n".join(f"- {c}" for c in ac)
            desc += f"\n\n[Created by SoftProdigy Snow Agent · module {s['module']}]"
            impact, urgency = _IU[s["priority"]]
            payload: Dict[str, Any] = {"short_description": s["title"][:160], "description": desc}
            if table in ("incident", "problem"):
                payload.update(impact=impact, urgency=urgency)
            elif table == "change_request":
                payload.update(type="normal", impact=impact, urgency=urgency,
                               justification=s["so_that"] or s["title"])
            elif table == "rm_story":
                payload.update(acceptance_criteria="\n".join(ac), story_points=s["story_points"],
                               priority={"critical": "1", "high": "2", "medium": "3", "low": "4"}[s["priority"]])
            else:
                payload["priority"] = {"critical": "1", "high": "2", "medium": "3", "low": "4"}[s["priority"]]
            # Only an explicit group (per story, an explicit module → group mapping, or one for the whole request) is
            # used, and only when it resolves to exactly one active group. Otherwise: no record, a question instead.
            mapping = body.get("group_mapping") if isinstance(body.get("group_mapping"), dict) else {}
            grp_name = str(raw.get("assignment_group") or mapping.get(s["module"]) or body.get("assignment_group")
                           or "").strip()
            group = None
            if grp_name:
                group, problem = _find_group(sn, grp_name, groups)
                if problem:
                    errors.append({"index": idx, "title": s["title"], "table": table, "needs_input": True,
                                   "error": problem + " Nothing was created for this story."})
                    continue
                payload["assignment_group"] = group["sys_id"]
            if table == "incident" and body.get("caller_id"):
                payload["caller_id"] = body["caller_id"]
            try:
                rec = sn.create(table, payload, "sys_id,number")
            except SNError as e:
                errors.append({"index": idx, "title": s["title"], "table": table, "error": str(e)})
                continue
            item = {"index": idx, "sys_id": rec.get("sys_id"), "number": rec.get("number"),
                    "table": table, "story_type": stype if stype in ("incident", "problem", "change")
                    else "story", "assignment_group": group["name"] if group else None,
                    "link": sn.link(table, rec.get("sys_id", ""))}
            if retyped:
                item["retyped_from"] = retyped
            # pipeline=true (the UI): a story that describes configuration is PLANNED (read-only: inspect, compare,
            # plan) and returned for review; nothing is built until the user confirms via /snow/requirements/deploy.
            if use_pipeline and stype not in ("incident", "problem", "change") and (
                    story_build.is_buildable(s) or rp.looks_like_configuration(story_build.story_text(s))):
                try:
                    item["requirement_plan"] = rp.plan_requirement(
                        sn, text=story_build.story_text(s),
                        story_ref={"table": table, "sys_id": rec.get("sys_id"), "number": rec.get("number")})
                except (RuntimeError, SNError) as e:
                    item["requirement_plan_error"] = str(e)
            # Stories that describe platform code also get that code built (script includes, business rules ...)
            elif build_artifacts and stype not in ("incident", "problem", "change") and story_build.is_buildable(s):
                build = story_build.build_story_artifacts(sn, s, body.get("update_set_name") or "")
                item["build"] = build
                try:
                    sn.patch(table, rec["sys_id"], {"work_notes": story_build.work_note(build)})
                except SNError as e:
                    logger.warning(f"could not add the build work note to {table} {rec.get('sys_id')}: {e}")
            created.append(item)
        builds = [c["build"] for c in created if c.get("build")]
        return {"success": not errors and all(b["status"] != "failed" for b in builds),
                "stories_created": created, "errors": errors,
            "planner_preflight": readiness,
                "summary": {"requested": len(stories), "created": len(created), "failed": len(errors),
                            "artifacts_built": sum(1 for b in builds for a in b["artifacts"]
                                                   if a["status"] in ("created", "reused")),
                            "builds_failed": sum(1 for b in builds if b["status"] == "failed"),
                            "requirements_planned": sum(1 for c in created if c.get("requirement_plan")),
                            "by_type": dict(Counter(c["story_type"] for c in created)),
                            "story_table": "rm_story" if has_story else "task"}}
    return await asyncio.to_thread(work)


# ════════════════════════════════════════════════════════════════════════════
# Health scan (UI format) and chat
# ════════════════════════════════════════════════════════════════════════════

@snow_universal_router.post("/snow/health-scan")
async def health_scan_ui(body: dict):
    sn = snc.client_from(body)
    return await asyncio.to_thread(sn_health.run_scan, sn)


@snow_universal_router.post("/snow/universal/chat")
async def universal_chat(body: dict):
    msg = (body.get("message") or "").strip()
    if not msg:
        raise HTTPException(status_code=400, detail="message is required")

    # Enrich context with release info when instance credentials are provided
    release_ctx = ""
    if body.get("snow_url") and body.get("snow_user") and body.get("snow_pass"):
        try:
            sn = snc.SNClient(body["snow_url"].rstrip("/"), body["snow_user"], body["snow_pass"])
            release = await asyncio.to_thread(sn.release_info)
            if release.get("build_name"):
                release_ctx = (f" Connected instance is running {release['build_name']}"
                               f" (tag: {release.get('build_tag') or 'n/a'}).")
        except Exception:
            pass

    system = ("You are SoftProdigy Snow, an expert ServiceNow consultant (ITSM, ITOM, CSM, HRSD, platform scripting, "
              "Flow Designer, update sets, ATF). Answer precisely, with short Markdown sections and code where "
              "useful. You cannot change the instance from chat." + release_ctx + " " + (body.get("context") or ""))
    msgs = [{"role": "system", "content": system}]
    for m in (body.get("conversation_history") or [])[-10:]:
        if isinstance(m, dict) and m.get("role") in ("user", "assistant") and m.get("content"):
            msgs.append({"role": m["role"], "content": str(m["content"])[:4000]})
    msgs.append({"role": "user", "content": msg})
    try:
        text = await asyncio.to_thread(groq_client.chat, msgs, 2000, 0.4)
    except groq_client.GroqError as e:
        raise HTTPException(status_code=503, detail=f"AI service unavailable: {e}")
    return {"success": True, "response": text}


# ════════════════════════════════════════════════════════════════════════════
# Jira
# ════════════════════════════════════════════════════════════════════════════

def _adf(text: str) -> dict:
    return {"type": "doc", "version": 1, "content": [
        {"type": "paragraph", "content": [{"type": "text", "text": line}]} for line in text.split("\n") if line.strip()
    ] or [{"type": "paragraph", "content": [{"type": "text", "text": " "}]}]}


@snow_universal_router.post("/snow/jira/push")
async def jira_push(body: dict):
    """
    Create Jira issues from stories. Uses JIRA_URL, JIRA_EMAIL, JIRA_API_TOKEN from .env (or the body)
    and body.project_key. Issue type defaults to Story.
    """
    url = (body.get("jira_url") or os.getenv("JIRA_URL", "")).rstrip("/")
    email = body.get("jira_email") or os.getenv("JIRA_EMAIL", "")
    token = body.get("jira_api_token") or os.getenv("JIRA_API_TOKEN", "")
    project = body.get("project_key") or os.getenv("JIRA_PROJECT_KEY", "")
    if not (url and email and token and project):
        raise HTTPException(status_code=400, detail="Jira is not configured: set JIRA_URL, JIRA_EMAIL, "
                                                    "JIRA_API_TOKEN and project_key")
    stories = [s for s in (_norm_story(x) for x in body.get("user_stories") or []) if s]
    if not stories:
        raise HTTPException(status_code=400, detail="No stories provided")

    def work():
        created, errors = [], []
        for i, s in enumerate(stories):
            desc = f"As a {s['as_a']}, I want {s['i_want']}" + (f", so that {s['so_that']}." if s["so_that"] else ".")
            if s["acceptance_criteria"]:
                desc += "\n\nAcceptance Criteria:\n" + "\n".join(f"- {c}" for c in s["acceptance_criteria"])
            payload = {"fields": {"project": {"key": project}, "summary": s["title"][:250],
                                  "issuetype": {"name": body.get("issue_type", "Story")},
                                  "description": _adf(desc), "labels": ["softprodigy", s["module"].lower()]}}
            try:
                r = requests.post(f"{url}/rest/api/3/issue", json=payload, auth=(email, token), timeout=30,
                                  headers={"Accept": "application/json"})
                if r.ok:
                    key = r.json().get("key")
                    created.append({"index": i, "key": key, "url": f"{url}/browse/{key}"})
                else:
                    errors.append({"index": i, "title": s["title"], "error": f"{r.status_code}: {r.text[:200]}"})
            except requests.RequestException as e:
                errors.append({"index": i, "title": s["title"], "error": str(e)[:200]})
        return {"success": not errors, "created": created, "errors": errors}
    return await asyncio.to_thread(work)
