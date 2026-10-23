"""
Shared Word (.docx) helpers for the SoftProdigy Snow Agent.

AI models answer in Markdown (**bold**, bullet lists, | tables |, ### headings).
Writing that text straight into a Word paragraph leaves literal asterisks and
pipes in the document. `add_markdown` converts it into real Word formatting.
"""

from __future__ import annotations

import re
from typing import Any, Iterable, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BRAND = RGBColor(0x1A, 0x73, 0xE8)
GREY = RGBColor(0x5F, 0x63, 0x68)

_INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|`[^`]+`)")


def s(value: Any) -> str:
    """Safe string for Word: None -> '', lists joined, numbers stringified."""
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return ", ".join(s(v) for v in value if v is not None)
    if isinstance(value, dict):
        return "; ".join(f"{k}: {s(v)}" for k, v in value.items())
    return str(value)


def as_list(value: Any) -> List[str]:
    """Normalise list-ish AI output (list, newline string, single string) to a list of strings."""
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        return [s(v).strip() for v in value if s(v).strip()]
    text = s(value).strip()
    if not text:
        return []
    parts = [re.sub(r"^\s*(?:[-*•]|\d+[.)])\s*", "", p).strip() for p in re.split(r"\n+", text)]
    return [p for p in parts if p]


def new_document() -> Document:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.1)
        section.right_margin = Inches(1.1)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    return doc


def heading(doc, text: str, level: int = 1, color: Optional[RGBColor] = BRAND):
    h = doc.add_heading(s(text) or " ", level=min(max(level, 0), 9))
    if color is not None:
        for r in h.runs:
            r.font.color.rgb = color
    return h


def title_page(doc, title: str, subtitle: str = "", meta: str = "", page_break: bool = True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    r = p.add_run(s(title))
    r.font.size = Pt(26)
    r.bold = True
    r.font.color.rgb = BRAND
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(s(subtitle))
        r2.font.size = Pt(16)
    if meta:
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(s(meta))
        r3.font.size = Pt(10)
        r3.font.color.rgb = GREY
    if page_break:
        doc.add_page_break()


def add_inline(paragraph, text: str, bold: bool = False, italic: bool = False):
    """Add text to a paragraph, converting **bold**, *italic* and `code` markers."""
    for part in _INLINE.split(s(text)):
        if not part:
            continue
        b, i, code = bold, italic, False
        if (part.startswith("**") and part.endswith("**")) or (part.startswith("__") and part.endswith("__")):
            part, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            part, code = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        run = paragraph.add_run(part)
        run.bold = b or None
        run.italic = i or None
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
    return paragraph


def shade(cell, fill: str = "E8F0FE"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def table(doc, headers: List[str], rows: Iterable[Iterable[Any]], widths: Optional[List[float]] = None):
    rows = [list(r) for r in rows]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        add_inline(cell.paragraphs[0], h, bold=True)
        shade(cell)
    for row in rows:
        cells = t.add_row().cells
        for i in range(len(headers)):
            val = row[i] if i < len(row) else ""
            cells[i].text = ""
            lines = as_list(val) if isinstance(val, (list, tuple)) else [s(val)]
            for j, line in enumerate(lines or [""]):
                para = cells[i].paragraphs[0] if j == 0 else cells[i].add_paragraph()
                add_inline(para, line)
                for r in para.runs:
                    r.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def _split_row(line: str) -> List[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip().replace("<br>", "\n").replace("<br/>", "\n") for c in line.split("|")]


def add_markdown(doc, text: str, base_level: int = 2):
    """Render a Markdown string into the document with real Word formatting."""
    lines = s(text).replace("\r\n", "\n").split("\n")
    i = 0
    para_buf: List[str] = []

    def flush():
        if para_buf:
            add_inline(doc.add_paragraph(), " ".join(x.strip() for x in para_buf))
            para_buf.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush(); i += 1; continue
        if stripped.startswith("```"):  # code block
            flush()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code))
            r.font.name = "Consolas"; r.font.size = Pt(9)
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            flush()
            headers = _split_row(stripped)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i])); i += 1
            table(doc, headers, rows)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            flush()
            heading(doc, m.group(2).strip("*# "), level=min(base_level + len(m.group(1)) - 1, 4), color=None)
            i += 1; continue
        m = re.match(r"^(\s*)[-*•+]\s+(.*)$", line)
        if m:
            flush()
            style = "List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet"
            add_inline(doc.add_paragraph(style=style), m.group(2)); i += 1; continue
        m = re.match(r"^\s*(\d+)[.)]\s+(.*)$", line)
        if m:
            flush()
            add_inline(doc.add_paragraph(style="List Number"), m.group(2)); i += 1; continue
        if re.match(r"^\s*(-{3,}|\*{3,})\s*$", line):
            flush(); i += 1; continue
        para_buf.append(line)
        i += 1
    flush()


def safe_filename(name: str, default: str = "document") -> str:
    out = "".join(c if c.isalnum() or c in ("_", "-") else "_" for c in s(name)).strip("_")
    return re.sub(r"_+", "_", out) or default
