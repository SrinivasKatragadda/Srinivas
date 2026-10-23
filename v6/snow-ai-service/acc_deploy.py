"""
ACC Deployment — ServiceNow Agent Client Collector (ACC-VC) global endpoint roll-out.

What the agent does with it (reference: itc_infotech__acc_deployment package):

1. GENERATE  functional spec / notes  ->  deployment package
             * <Prefix>_Technical_Specification.docx  (15 sections + appendix, like the reference TS)
             * <Prefix>_Architecture.puml / .svg / .png
             * <Prefix>_WavePlan.puml / .svg / .png
             * <Prefix>_program.json   (waves, gates, brand tiers — used to seed ServiceNow)
2. SET UP    the program in ServiceNow, inside an application scope + update set (reuses the Build App engine):
             tables acc_wave, acc_stage_gate, acc_gate_decision (audit, immutable), acc_brand_health, acc_deploy_error,
             tracking fields on cmdb_ci_computer, roles, and business rules that block GO without a human approver.
             Then seeds waves / gates / brand tiers as records.
3. COMMAND CENTER (live, read-only)  KPIs from the instance: coverage vs target, registration success, freshness,
             duplicate CIs, completeness, OS / region / brand splits, reconciliation alerts, AI remediation briefing.
4. WAVE PLANNER  waves (status / deployed counts, reconciled against the target) and stage gates. A gate decision
             is only recorded with the signed-in user as approver, a reason of at least 50 characters, and an
             audit row. The AI gives an advisory recommendation only — it can never approve a gate.
"""

from __future__ import annotations

import asyncio
import base64
import calendar
import io
import json
import logging
import os
import re
import textwrap
import threading
import time
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import JSONResponse, Response

import arch_diagram
import docx_render as dr
import groq_client
import sn_build
import sn_client as snc
from sn_client import SNError

logger = logging.getLogger(__name__)
acc_router = APIRouter()

PKG_DIR = Path(os.getenv("AGENT_JOBS_DIR", Path(__file__).resolve().parent / "jobs")) / "acc"
MAX_SOURCE = 60000

LOGICAL_TABLES = ["acc_wave", "acc_stage_gate", "acc_gate_decision", "acc_brand_health", "acc_deploy_error"]
CMDB_FIELDS = ["region_code", "wave_id", "brand_tier", "brand_name", "acc_vc_status", "intune_enrolled",
               "duplicate_risk", "deploy_confidence"]
CI_TABLE = "cmdb_ci_computer"
ACC_AGENT_TABLE = "sn_agent_cmdb_ci_agent"   # ServiceNow ACC plugin (Agent Client Collector) CI table

# ════════════════════════════════════════════════════════════════════════════
# Program model
# ════════════════════════════════════════════════════════════════════════════

DEFAULT_PROGRAM: Dict[str, Any] = {
    "title": "ACC-VC MID-less Global Endpoint Deployment",
    "client": "Client",
    "delivery_partner": "Delivery Partner",
    "product": "ServiceNow ACC-VC (Agent Client Collector – Visibility & Configuration)",
    "prefix": "ACC_VC",
    "deployment_channel": "Microsoft Intune",
    "total_endpoints": 250000,
    "timeline": "Jun 2026 – Jul 2027",
    "os_split": [{"os": "Windows", "count": 125000}, {"os": "macOS", "count": 75000}, {"os": "Linux", "count": 50000}],
    "regions": [{"code": "NA", "name": "North America"}, {"code": "EMEA", "name": "EMEA"},
                {"code": "APAC", "name": "APAC"}, {"code": "CALA", "name": "CALA"},
                {"code": "GLOBAL", "name": "Global / Corporate"}],
    "brand_tiers": [
        {"tier": "LUXURY", "brands": 7, "target": 28400, "health": "91.2% healthy", "issues": "Firewall blocks in two regions"},
        {"tier": "PREMIUM", "brands": 11, "target": 89200, "health": "79.9% healthy", "issues": "Two site clusters flagged"},
        {"tier": "SELECT", "brands": 10, "target": 94600, "health": "68.7% healthy", "issues": "Sites not installed; proxy issue"},
        {"tier": "CORPORATE", "brands": 0, "target": 37800, "health": "[TBD – baseline required]", "issues": "[TBD]"}],
    "waves": [
        {"wave_id": "WAVE-0", "name": "Pilot Foundation", "start": "2026-06", "end": "2026-07",
         "regions": "All (controlled groups)", "target": 500,
         "entry": "Intune enrollment validated for pilot group; Gate 1 human approval obtained; network connectivity confirmed",
         "exit": "Pilot health > 90%; 0 duplicate CIs; 3-day KPI stability; Gate 2 human approval"},
        {"wave_id": "WAVE-1", "name": "North America Controlled Rollout", "start": "2026-08", "end": "2026-09",
         "regions": "North America", "target": 45000,
         "entry": "Wave 0 exit criteria met; Gate 2 approved; NA Intune groups configured",
         "exit": "LUXURY ≥ 90% health; 3-day KPI stability; Gate 3 inputs ready"},
        {"wave_id": "WAVE-2", "name": "EMEA Expansion", "start": "2026-10", "end": "2026-11", "regions": "EMEA",
         "target": 55000, "entry": "Gate 3 approved; EMEA firewall/proxy remediation completed",
         "exit": "PREMIUM ≥ 85% health for EMEA slice; duplicate CI risk < 5%"},
        {"wave_id": "WAVE-3", "name": "APAC Expansion", "start": "2026-12", "end": "2027-01", "regions": "APAC",
         "target": 50000, "entry": "EMEA exit criteria met; APAC proxy fix confirmed",
         "exit": "APAC slice health within wave targets; freshness SLA ≥ 90%"},
        {"wave_id": "WAVE-4", "name": "CALA & Global Functions", "start": "2027-02", "end": "2027-04",
         "regions": "CALA + Global/Corporate", "target": 62700,
         "entry": "Wave 3 exit; blocked sites unblocked; Corporate Intune groups ready",
         "exit": "CALA + Corporate deployed; SELECT ≥ 75% health; cumulative total ≥ 213,200"},
        {"wave_id": "WAVE-5", "name": "Global Stabilization & Optimization", "start": "2027-05", "end": "2027-07",
         "regions": "All remaining + remediation", "target": 36800,
         "entry": "All prior wave exit criteria met; Gate 4 human approval",
         "exit": "250,000 total reconciled; duplicate CI count = 0; freshness ≥ 95%; executive sign-off"}],
    "gates": [
        {"gate_id": "GATE-1", "name": "Architecture, Security & Platform Readiness", "unlocks_wave": "WAVE-0",
         "objective": "Confirm MID-less architecture, mTLS certificates, Intune connectivity and instance readiness",
         "owner": "Delivery Lead", "approver": "Platform Owner + Security Lead",
         "criteria": "Architecture review complete; security sign-off; Intune enrollment > 0 for pilot group"},
        {"gate_id": "GATE-2", "name": "Pilot Deployment Validation", "unlocks_wave": "WAVE-1",
         "objective": "Confirm the Wave 0 pilot met all health targets before scale-out",
         "owner": "Endpoint Engineering Lead", "approver": "Regional Deployment Lead (NA)",
         "criteria": "Pilot health > 90%; 0 unsolved duplicate CIs; rollback tested successfully"},
        {"gate_id": "GATE-3", "name": "Regional Deployment Readiness", "unlocks_wave": "WAVE-2",
         "objective": "Confirm regional network and compliance requirements are resolved",
         "owner": "Regional Leads (EMEA + APAC)", "approver": "Delivery Leadership + Security",
         "criteria": "All regional network fixes confirmed; LUXURY ≥ 90%; PREMIUM ≥ 85%"},
        {"gate_id": "GATE-4", "name": "Global Production Readiness & Stabilization", "unlocks_wave": "WAVE-5",
         "objective": "Final sign-off that the endpoint target is met and data quality is stable",
         "owner": "Program Lead", "approver": "Executive Stakeholders",
         "criteria": "Target reconciled; freshness ≥ 95%; duplicate CI = 0; 3-day KPI stability"}],
    "architecture": {
        "title": "ACC-VC MID-less Architecture",
        "layers": [
            {"name": "Endpoint Estate", "components": ["Windows Endpoints", "macOS Endpoints", "Linux Endpoints"]},
            {"name": "Microsoft Intune", "components": ["Enrollment & Policy Engine", "Deployment Groups", "Win32 / Shell Packages"]},
            {"name": "ACC-VC Agent", "components": ["Agent Process", "mTLS Certificate Store", "Local Log Store"]},
            {"name": "Cloud Ingestion", "components": ["Ingestion Endpoint (mTLS)", "API + Registration Key Validator", "CI Registration Engine"]},
            {"name": "ServiceNow Platform", "components": ["CMDB cmdb_ci_computer", "Duplicate CI Detection", "Data Freshness Monitor", "Wave & Gate Governance"]},
            {"name": "Command Center", "components": ["Executive KPI Gauges", "Regional Breakdown", "AI Remediation Briefing", "Wave Planner"]}],
        "flows": [
            {"from": "Deployment Groups", "to": "Win32 / Shell Packages", "label": "deploy trigger"},
            {"from": "Agent Process", "to": "mTLS Certificate Store", "label": "client cert"},
            {"from": "mTLS Certificate Store", "to": "Ingestion Endpoint (mTLS)", "label": "HTTPS 443 + mTLS"},
            {"from": "Ingestion Endpoint (mTLS)", "to": "API + Registration Key Validator", "label": "validate keys"},
            {"from": "CI Registration Engine", "to": "CMDB cmdb_ci_computer", "label": "create/update CI"},
            {"from": "Duplicate CI Detection", "to": "AI Remediation Briefing", "label": "risk alerts"},
            {"from": "Wave & Gate Governance", "to": "Wave Planner", "label": "gate status"}]},
}


def _int(v: Any) -> Optional[int]:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v)
    m = re.search(r"-?\d[\d,]*", str(v).replace("~", ""))
    if not m:
        return None
    try:
        return int(m.group(0).replace(",", ""))
    except ValueError:
        return None


def _txt(v: Any, n: int = 4000) -> str:
    if v is None:
        return ""
    if isinstance(v, (list, tuple)):
        return "; ".join(_txt(x, n) for x in v if x not in (None, ""))[:n]
    return str(v).strip()[:n]


_MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_abbr) if m}


def parse_month(v: Any, end: bool = False) -> Optional[date]:
    """'2026-06', '2026-06-15', 'Jun 2026', 'June 2026' -> date (1st of month, or last day when end=True)."""
    s = str(v or "").strip()
    m = re.match(r"^(\d{4})-(\d{1,2})(?:-(\d{1,2}))?", s)
    if m:
        y, mo = int(m.group(1)), int(m.group(2))
        if m.group(3):
            try:
                return date(y, mo, int(m.group(3)))
            except ValueError:
                pass
    else:
        m = re.match(r"^([A-Za-z]{3})[A-Za-z]*\.?\s+(\d{4})", s)
        if not m or m.group(1).lower() not in _MONTHS:
            return None
        y, mo = int(m.group(2)), _MONTHS[m.group(1).lower()]
    if not 1 <= mo <= 12:
        return None
    return date(y, mo, calendar.monthrange(y, mo)[1] if end else 1)


def _slugid(v: str, prefix: str) -> str:
    s = re.sub(r"[^A-Z0-9]+", "-", str(v or "").upper()).strip("-")
    return s if s.startswith(prefix) else f"{prefix}-{s}" if s else prefix


def normalize_program(p: Any) -> Dict[str, Any]:
    p = p if isinstance(p, dict) else {}
    out: Dict[str, Any] = {
        "title": _txt(p.get("title"), 200) or DEFAULT_PROGRAM["title"],
        "client": _txt(p.get("client"), 120) or "[TBD – client]",
        "delivery_partner": _txt(p.get("delivery_partner"), 120),
        "product": _txt(p.get("product"), 200) or DEFAULT_PROGRAM["product"],
        "deployment_channel": _txt(p.get("deployment_channel"), 120) or "Microsoft Intune",
        "total_endpoints": _int(p.get("total_endpoints")),
        "timeline": _txt(p.get("timeline"), 80),
    }
    pre = re.sub(r"[^A-Za-z0-9]+", "_", _txt(p.get("prefix"), 40) or "ACC_VC").strip("_") or "ACC_VC"
    out["prefix"] = pre[:40]
    out["os_split"] = [{"os": _txt(o.get("os"), 40), "count": _int(o.get("count"))}
                       for o in (p.get("os_split") or []) if isinstance(o, dict) and o.get("os")][:8]
    out["regions"] = [{"code": _txt(r.get("code") or r.get("name"), 20).upper(), "name": _txt(r.get("name") or r.get("code"), 60)}
                      for r in (p.get("regions") or []) if isinstance(r, dict) and (r.get("code") or r.get("name"))][:12]
    out["brand_tiers"] = [{"tier": _txt(b.get("tier"), 40).upper(), "brands": _int(b.get("brands")),
                           "target": _int(b.get("target")), "health": _txt(b.get("health"), 80),
                           "issues": _txt(b.get("issues"), 400)}
                          for b in (p.get("brand_tiers") or []) if isinstance(b, dict) and b.get("tier")][:10]
    waves = []
    for i, w in enumerate(p.get("waves") or []):
        if not isinstance(w, dict):
            continue
        wid = _slugid(w.get("wave_id") or f"WAVE-{i}", "WAVE")
        waves.append({"wave_id": wid, "name": _txt(w.get("name"), 100) or wid,
                      "start": _txt(w.get("start"), 20), "end": _txt(w.get("end"), 20),
                      "regions": _txt(w.get("regions"), 200), "target": _int(w.get("target")),
                      "entry": _txt(w.get("entry"), 1500), "exit": _txt(w.get("exit"), 1500)})
    out["waves"] = waves[:12]
    gates = []
    wave_ids = [w["wave_id"] for w in out["waves"]]
    for i, g in enumerate(p.get("gates") or []):
        if not isinstance(g, dict):
            continue
        gid = _slugid(g.get("gate_id") or f"GATE-{i + 1}", "GATE")
        unlocks = _slugid(g.get("unlocks_wave"), "WAVE") if g.get("unlocks_wave") else ""
        gates.append({"gate_id": gid, "name": _txt(g.get("name"), 200) or gid,
                      "objective": _txt(g.get("objective"), 1500), "owner": _txt(g.get("owner"), 200),
                      "approver": _txt(g.get("approver"), 200), "criteria": _txt(g.get("criteria"), 1500),
                      "unlocks_wave": unlocks if unlocks in wave_ids else ""})
    out["gates"] = gates[:10]
    arch = arch_diagram.normalize_spec(p.get("architecture") or {})
    out["architecture"] = arch if arch["layers"] else arch_diagram.normalize_spec(DEFAULT_PROGRAM["architecture"])
    return out


def reconcile(program: Dict[str, Any]) -> List[Dict[str, str]]:
    """Every breakdown must add up to the total endpoint target (reference TS §4.3 rule)."""
    out: List[Dict[str, str]] = []
    total = program.get("total_endpoints")
    if not total:
        out.append({"level": "warning", "message": "No total endpoint target — reconciliation cannot be checked."})
        return out
    for label, items, key in (("Brand tiers", program.get("brand_tiers"), "target"),
                              ("Waves", program.get("waves"), "target"),
                              ("OS split", program.get("os_split"), "count")):
        vals = [i.get(key) for i in items or []]
        if not vals:
            continue
        known = [v for v in vals if isinstance(v, int)]
        s_ = sum(known)
        if len(known) < len(vals):
            out.append({"level": "warning", "message": f"{label}: {len(vals) - len(known)} item(s) without a number — "
                                                       f"known total {s_:,} of {total:,}."})
        elif s_ != total:
            out.append({"level": "error", "message": f"{label} add up to {s_:,}, not {total:,} "
                                                     f"(delta {total - s_:+,}). Reconcile before wave progression."})
        else:
            out.append({"level": "ok", "message": f"{label} reconcile to exactly {total:,}."})
    return out


# ════════════════════════════════════════════════════════════════════════════
# Diagrams — wave plan (Gantt + gates) and PlantUML sources
# ════════════════════════════════════════════════════════════════════════════

def render_waveplan(program: Dict[str, Any], fmt: str = "png", dpi: int = 180) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.dates as mdates
    import matplotlib.pyplot as plt

    waves = [w for w in program.get("waves") or [] if parse_month(w.get("start")) and parse_month(w.get("end"), True)]
    if not waves:
        raise ValueError("no waves with start/end months")
    n = len(waves)
    fig, ax = plt.subplots(figsize=(11, 1.1 + 0.62 * n))
    colors = ["#0F766E", "#1E5EBF", "#7C3AED", "#B45309", "#BE123C", "#15803D", "#475569"]
    cum = 0
    for i, w in enumerate(waves):
        s, e = parse_month(w["start"]), parse_month(w["end"], True) + timedelta(days=1)
        y = n - 1 - i
        ax.barh(y, (e - s).days, left=mdates.date2num(s), height=0.52, color=colors[i % len(colors)], alpha=0.9)
        tgt = w.get("target")
        cum += tgt or 0
        label = f"{w['wave_id']}  {w['name']}"
        ax.text(mdates.date2num(s) - 4, y, label, ha="right", va="center", fontsize=8, color="#0A1628")
        ax.text(mdates.date2num(e) + 4, y, (f"~{tgt:,}  (cum. {cum:,})" if tgt else ""), ha="left", va="center",
                fontsize=7.5, color="#334155")
    for g in program.get("gates") or []:
        idx = next((i for i, w in enumerate(waves) if w["wave_id"] == g.get("unlocks_wave")), None)
        if idx is None:
            continue
        x = mdates.date2num(parse_month(waves[idx]["start"]))
        y = n - 1 - idx + 0.42
        ax.plot([x], [y], marker="D", markersize=9, color="#DC2626", zorder=5)
        ax.text(x + 3, y + 0.02, g["gate_id"].replace("GATE-", "G"), fontsize=7.5, color="#DC2626",
                fontweight="bold", va="center")
    ax.xaxis_date()
    ax.xaxis.set_major_locator(mdates.MonthLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b\n%Y"))
    ax.tick_params(axis="x", labelsize=7)
    ax.set_yticks([])
    for sp in ("top", "right", "left"):
        ax.spines[sp].set_visible(False)
    ax.grid(axis="x", color="#E2E8F0", lw=0.6)
    first = parse_month(waves[0]["start"])
    last = parse_month(waves[-1]["end"], True)
    span = (last - first).days
    ax.set_xlim(mdates.date2num(first) - span * 0.42, mdates.date2num(last) + span * 0.2)
    ax.set_ylim(-0.7, n - 0.1)
    ax.set_title(f"{program.get('title', 'Deployment')} — {n}-Wave Rollout & Stage Gates", fontsize=10.5,
                 fontweight="bold", color="#0A1628", loc="left")
    ax.text(0, -0.16, "◆ = stage gate that must be approved by a named human before the wave starts (no auto-approval)",
            transform=ax.transAxes, fontsize=7.5, color="#DC2626")
    buf = io.BytesIO()
    fig.savefig(buf, format=fmt, dpi=dpi, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def _puml_text(s: str) -> str:
    return str(s or "").replace('"', "'").replace("\n", " ")


def architecture_puml(program: Dict[str, Any]) -> str:
    spec = program["architecture"]
    tints = ["#EFF3FB", "#E8F5E9", "#FFF3E0", "#FCE4EC", "#E8EAF6", "#EDE7F6"]
    lines = [f"@startuml {program['prefix']}_Architecture", "skinparam backgroundColor #FFFFFF",
             "skinparam componentStyle rectangle", "skinparam defaultFontName Arial", "skinparam defaultFontSize 11",
             "skinparam ArrowColor #555555", "skinparam PackageBorderColor #3F51B5", "",
             f"title {_puml_text(spec['title'])}", ""]
    alias: Dict[str, str] = {}
    k = 0
    for i, layer in enumerate(spec["layers"]):
        lines.append(f'package "{_puml_text(layer["name"])}" as L{i} {tints[i % len(tints)]} {{')
        for c in layer["components"]:
            k += 1
            alias[c] = f"C{k}"
            lines.append(f'  [{_puml_text(c)}] as C{k}')
        lines.append("}")
        lines.append("")
    for f in spec["flows"]:
        lines.append(f"{alias[f['from']]} --> {alias[f['to']]}" + (f" : {_puml_text(f['label'])}" if f["label"] else ""))
    lines.append("@enduml")
    return "\n".join(lines)


def waveplan_puml(program: Dict[str, Any]) -> str:
    lines = [f"@startuml {program['prefix']}_WavePlan", "skinparam backgroundColor #FFFFFF",
             "skinparam defaultFontName Arial", "skinparam defaultFontSize 11",
             "skinparam ActivityBackgroundColor #EFF3FB", "skinparam ActivityBorderColor #3F51B5",
             f"title {_puml_text(program['title'])} – {len(program['waves'])}-Wave Rollout & Stage Gate Governance"
             + (f" ({_puml_text(program['timeline'])})" if program.get("timeline") else ""), "", "start"]
    gates_for = {}
    for g in program["gates"]:
        gates_for.setdefault(g.get("unlocks_wave") or "", []).append(g)
    for w in program["waves"]:
        for g in gates_for.get(w["wave_id"], []):
            lines += [f"|#FCE4EC|{_puml_text(g['gate_id'])}|",
                      f":{_puml_text(g['gate_id'])} – {_puml_text(g['name'])}\\n**Human Approval Required**;",
                      f"note right: Approver: {_puml_text(g.get('approver') or '[TBD]')}"]
        tgt = f"~{w['target']:,} endpoints" if w.get("target") else "[TBD] endpoints"
        lines += [f"|#E8F5E9|{_puml_text(w['wave_id'])} – {_puml_text(w['name'])}|",
                  f":{_puml_text(w['wave_id'])} – {_puml_text(w['name'])}\\n{_puml_text(w.get('start'))} – "
                  f"{_puml_text(w.get('end'))}\\n{tgt};"]
        if w.get("exit"):
            lines.append(f"note right: Exit: {_puml_text(w['exit'])[:160]}")
    for g in gates_for.get("", []):
        lines += [f"|#FCE4EC|{_puml_text(g['gate_id'])}|",
                  f":{_puml_text(g['gate_id'])} – {_puml_text(g['name'])}\\n**Human Approval Required**;"]
    total = program.get("total_endpoints")
    lines += [f":DEPLOYMENT COMPLETE\\n{total:,} Endpoints Reconciled;" if total else ":DEPLOYMENT COMPLETE;",
              "stop", "@enduml"]
    return "\n".join(lines)


# ════════════════════════════════════════════════════════════════════════════
# AI — program extraction + Technical Specification sections
# ════════════════════════════════════════════════════════════════════════════

SYSTEM = ("You are a senior ServiceNow ACC (Agent Client Collector) / CMDB deployment architect writing an "
          "implementation-ready Technical Specification for an enterprise endpoint roll-out. Be specific to the "
          "source material. Never invent real API keys, registration keys, certificates or URLs — always use "
          "placeholders such as <SERVICENOW_INSTANCE_URL>, <ACC_VC_API_KEY>, <ACC_VC_REGISTRATION_KEY>, <REGION_CODE>, "
          "<WAVE_ID>, <PROXY_URL_IF_REQUIRED>. Anything the source does not state is written as "
          "'[TBD – <what is needed>]' — never guessed as fact.")

PROGRAM_PROMPT = """ACC_PROGRAM_EXTRACT
Read the source material and extract the deployment program as JSON with exactly these keys:
{
 "title": "short program title",
 "client": "client organisation", "delivery_partner": "implementation partner or ''",
 "product": "product being deployed", "prefix": "short file prefix like ACC_VC",
 "deployment_channel": "e.g. Microsoft Intune", "timeline": "e.g. Jun 2026 – Jul 2027",
 "total_endpoints": integer or null,
 "os_split": [{"os": "Windows", "count": integer or null}],
 "regions": [{"code": "NA", "name": "North America"}],
 "brand_tiers": [{"tier": "LUXURY", "brands": integer or null, "target": integer or null, "health": "text", "issues": "text"}],
 "waves": [{"wave_id": "WAVE-0", "name": "...", "start": "YYYY-MM", "end": "YYYY-MM", "regions": "...",
            "target": integer or null, "entry": "entry criteria", "exit": "exit criteria"}],
 "gates": [{"gate_id": "GATE-1", "name": "...", "objective": "...", "owner": "...", "approver": "...",
            "criteria": "...", "unlocks_wave": "the WAVE-n that may only start after this gate is approved"}],
 "architecture": {"title": "...", "layers": [{"name": "layer", "components": ["2-5 short names"]}],
                  "flows": [{"from": "component name", "to": "component name", "label": "max 3 words"}]}
}
Rules: numbers exactly as in the source (no rounding, no invention; null when absent). 3-6 architecture layers,
max 5 components each, max 8 flows whose from/to are component names used in layers. If the source has no wave
plan or gates, propose a sensible phased plan and mark names/dates with '[Proposed]'.

SOURCE:
<<<
{source}
>>>"""

SECTION_PROMPT = """ACC_TS_SECTIONS
Write these sections of the Technical Specification. Program facts (authoritative — use these numbers):
{program}

Source material (excerpt):
<<<
{source}
>>>

SECTIONS TO WRITE (use exactly these numbers and headings, in this order):
{outline}

Return JSON: {{"sections": [{{"heading": "4.1 Requirements Traceability Table", "level": 2,
  "blocks": [ {{"type": "paragraph", "text": "..."}},
              {{"type": "bullets", "items": ["..."]}},
              {{"type": "numbered", "items": ["..."]}},
              {{"type": "table", "headers": ["..."], "rows": [["...", "..."]]}},
              {{"type": "code", "text": "multi-line commands or pseudocode"}},
              {{"type": "note", "text": "a rule or security note"}} ] }}]}}
level: 1 for 'N. Title', 2 for 'N.N Title', 3 for 'N.N.N Title'. Tables must be complete (every row filled,
no '...' rows). Be concrete: real ServiceNow table/field names, real commands with placeholders, measurable
thresholds."""

OUTLINES: List[Tuple[str, str]] = [
    ("intro", "2. Introduction (level 1, no blocks)\n2.1 Purpose — one paragraph\n"
              "2.2 Scope — paragraph 'In Scope:' + bullets, paragraph 'Out of Scope:' + bullets\n"
              "2.3 References — bullets (source document, the architecture diagram file "
              "{prefix}_Architecture.svg, the wave plan diagram file {prefix}_WavePlan.svg)"),
    ("arch", "3.2 Key Architectural Decisions — table Decision | Design (6-8 rows: MID Server usage, deployment "
             "channel, authentication, communication direction, credential model, CI source of truth, …)\n"
             "3.3 High-Level Runtime Flow — numbered list of 10-14 steps from device targeting to audit logging"),
    ("req", "4.1 Requirements Traceability Table — table Req ID | Source Section | Technical Module/File | "
            "Logic/Design Notes | Status (12-16 rows REQ-001…)\n4.2 Data Model (level 2, no blocks)\n"
            "4.2.1 Core CMDB Entities — table Entity | Field | Type | Key/Index | Notes (cmdb_ci_computer incl. "
            "u_region_code, u_wave_id, u_brand_tier, u_brand_name, u_acc_vc_status, u_intune_enrolled, "
            "u_duplicate_risk, u_deploy_confidence, serial_number, last_discovered)\n"
            "4.2.2 Wave Governance Entity — table (u_acc_wave)\n4.2.3 Stage Gate Entity — table (u_acc_stage_gate "
            "incl. human_approval_required always true, approver_identity, decision_reason)\n"
            "4.2.4 Brand-Tier Health Entity — table (u_acc_brand_health)"),
    ("logic", "4.4 Business Logic – Duplicate CI Detection — code block pseudocode (serial / hostname / MAC match, "
              "stale flag)\n4.5 Validations — table Field / Condition | Validation Logic | Error Message (8-10 rows "
              "incl. serial uniqueness, Intune enrollment, registration timeout, freshness, wave sum ≤ target, "
              "no auto-approval, decision reason ≥ 50 chars, proxy, mTLS certificate)\n"
              "4.6 Integrations — table System | Method/Protocol | Data Exchanged | Trigger/Frequency | Error Handling"),
    ("os", "5. OS-Specific Deployment Command Matrix (level 1) — note block: all keys/URLs are placeholders injected "
           "at runtime from a secure parameter store\n"
           "5.1 Windows – Win32 App Package — table Parameter | Value (package type, silent flags, log path) then "
           "code blocks: silent install, validation, uninstall, rollback\n"
           "5.2 macOS – Shell Script Package — same structure\n5.3 Linux – Shell Script Package — same structure "
           "(rpm and apt variants)\n5.4 Deployment Group Naming Convention — code block with pattern and 4 examples"),
    ("ops", "8. Pre-Deployment Checklist — table # | Category | Check Item | Pass Criteria | Owner (12-15 rows)\n"
            "9. Rollback Plan (level 1, no blocks)\n9.1 Rollback Triggers — table Trigger | Threshold | Action\n"
            "9.2 Rollback Steps by Scope — table Scope | Steps (Device-Level, Intune Policy, CI Cleanup, "
            "Post-Rollback Validation)\n9.3 Escalation Matrix — table Level | Trigger | Escalation Path (L1-L4)"),
    ("ui", "10. Command Center Design Specification (P1) (level 1, no blocks)\n10.1 KPI Gauge Panel — table KPI Gauge "
           "| Data Source | Healthy Threshold (8 gauges)\n10.2 Regional Breakdown Table — table Column | Description\n"
           "10.3 Wave Burndown Charts — bullets\n10.4 AI Remediation Briefing Panel — bullets\n"
           "10.5 Context Menu Actions — table # | Action | Target (12-16 rows)\n"
           "11. Wave Deployment Planner Design Specification (P2) (level 1, no blocks)\n11.1 Gantt View – Wave Bars — "
           "bullets\n11.2 Stage Gate Tracker — bullets\n11.3 Risk Heat Bar — table Risk Category | Score Source\n"
           "11.4 AI Risk Assessment Panel — bullets (advisory only, human approval required)"),
    ("close", "12. Non-Functional Considerations — table Concern | Specification (scalability, security, compliance, "
              "logging, monitoring, audit trail)\n13. Testing & Acceptance — table Test Type | Scope | Acceptance "
              "Criterion | Maps To (10-12 rows)\n14. Traceability Matrix — table Source (FS Section) | TS Section | "
              "Requirement ID\n15. Open Items — table # | Item | Section | Status (everything marked TBD in the "
              "source or in this spec)\nAppendix A: Deployment Readiness Checklist Summary (level 1) — table # | "
              "Checklist Item | Owner | Status (Status column '[ ]')"),
]

_gen_jobs: Dict[str, Dict[str, Any]] = {}
_gen_lock = threading.Lock()


def _source_excerpt(text: str, n: int) -> str:
    text = text or ""
    return text if len(text) <= n else text[:n] + "\n[… source truncated …]"


def _norm_sections(data: Any) -> List[Dict[str, Any]]:
    secs = (data or {}).get("sections") if isinstance(data, dict) else data
    out = []
    for s_ in secs or []:
        if not isinstance(s_, dict) or not s_.get("heading"):
            continue
        blocks = []
        for b in s_.get("blocks") or []:
            if isinstance(b, str):
                b = {"type": "paragraph", "text": b}
            if not isinstance(b, dict):
                continue
            t = str(b.get("type") or "paragraph").lower()
            if t in ("bullets", "numbered"):
                items = dr.as_list(b.get("items"))
                if items:
                    blocks.append({"type": t, "items": items})
            elif t == "table":
                headers = [dr.s(h) for h in (b.get("headers") or [])]
                rows = [[dr.s(c) for c in r] if isinstance(r, (list, tuple)) else
                        [dr.s(r.get(h, "")) for h in headers] if isinstance(r, dict) else [dr.s(r)]
                        for r in (b.get("rows") or [])]
                if headers:
                    blocks.append({"type": "table", "headers": headers, "rows": rows})
            elif t in ("code", "note", "paragraph"):
                if dr.s(b.get("text")).strip():
                    blocks.append({"type": t, "text": dr.s(b.get("text"))})
        lvl = s_.get("level")
        if not isinstance(lvl, int):
            lvl = str(s_["heading"]).split(" ")[0].count(".") + 1 if re.match(r"^\d", str(s_["heading"])) else 1
            if re.match(r"^\d+\.\s", str(s_["heading"])):
                lvl = 1
        out.append({"heading": dr.s(s_["heading"]).strip(), "level": max(1, min(3, lvl)), "blocks": blocks})
    return out


def _ai_sections(key: str, outline: str, program: Dict[str, Any], source: str) -> List[Dict[str, Any]]:
    prog = {k: v for k, v in program.items() if k != "architecture"}
    prompt = SECTION_PROMPT.format(program=json.dumps(prog, ensure_ascii=False)[:9000],
                                   source=_source_excerpt(source, 12000),
                                   outline=outline.replace("{prefix}", program["prefix"]))
    last: Optional[Exception] = None
    for attempt in range(2):
        try:
            data = groq_client.generate_json(prompt, system=SYSTEM, max_tokens=7000)
            secs = _norm_sections(data)
            if secs:
                return secs
            last = ValueError("the AI answer contained no sections")
        except Exception as e:  # retried once, then reported
            last = e
        time.sleep(1.5 * (attempt + 1))
    raise RuntimeError(f"section group '{key}' could not be generated: {last}")


def generate_package(job: Dict[str, Any], source: str, overrides: Dict[str, Any]):
    """Background job: program -> sections -> diagrams -> docx -> zip."""
    def step(msg: str):
        job["progress"].append({"time": datetime.now().strftime("%H:%M:%S"), "msg": msg})

    try:
        job["status"] = "running"
        step("Reading the deployment program from the source (waves, gates, brand tiers, architecture)…")
        if source.strip():
            try:
                raw = groq_client.generate_json(PROGRAM_PROMPT.replace("{source}", _source_excerpt(source, 30000)),
                                                system=SYSTEM, max_tokens=6000)
            except Exception as e:
                raise RuntimeError(f"AI could not read the program: {e}")
        else:
            raw = json.loads(json.dumps(DEFAULT_PROGRAM))
            step("No source given — using the reference ACC-VC program template.")
        for k, v in overrides.items():
            if v not in (None, ""):
                raw[k] = v
        program = normalize_program(raw)
        job["program"] = program
        job["reconciliation"] = reconcile(program)
        step(f"Program: {program['title']} · {len(program['waves'])} waves · {len(program['gates'])} gates · "
             f"target {program['total_endpoints'] or '[TBD]'} endpoints")

        results: Dict[str, List[Dict[str, Any]]] = {}
        errors: List[str] = []
        workers = int(os.getenv("ACC_AI_CONCURRENCY", "2"))
        step(f"Writing {len(OUTLINES)} section groups of the Technical Specification…")
        with ThreadPoolExecutor(max_workers=max(1, workers)) as ex:
            futs = {ex.submit(_ai_sections, k, o, program, source): k for k, o in OUTLINES}
            for f in futs:
                k = futs[f]
                try:
                    results[k] = f.result()
                    step(f"✓ {results[k][0]['heading'] if results[k] else k}")
                except Exception as e:
                    errors.append(str(e))
                    step(f"✗ {e}")
        if errors:
            raise RuntimeError("; ".join(errors))

        step("Drawing the architecture and wave-plan diagrams…")
        files: Dict[str, bytes] = {}
        pre = program["prefix"]
        arch_png = arch_diagram.render_png(program["architecture"])
        files[f"{pre}_Architecture.png"] = arch_png
        files[f"{pre}_Architecture.svg"] = arch_diagram.render_png(program["architecture"], fmt="svg")
        files[f"{pre}_Architecture.puml"] = architecture_puml(program).encode()
        wave_png = None
        try:
            wave_png = render_waveplan(program)
            files[f"{pre}_WavePlan.png"] = wave_png
            files[f"{pre}_WavePlan.svg"] = render_waveplan(program, fmt="svg")
        except ValueError as e:
            step(f"⚠ Wave plan chart skipped: {e}")
        files[f"{pre}_WavePlan.puml"] = waveplan_puml(program).encode()

        step("Writing the Word document…")
        docx = build_docx(program, results, arch_png, wave_png, job["reconciliation"])
        files[f"{pre}_Technical_Specification.docx"] = docx
        files[f"{pre}_program.json"] = json.dumps(program, indent=2, ensure_ascii=False).encode()

        PKG_DIR.mkdir(parents=True, exist_ok=True)
        zpath = PKG_DIR / f"{job['job_id']}.zip"
        with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
            for name, data in files.items():
                z.writestr(name, data)
        (PKG_DIR / f"{job['job_id']}.json").write_text(json.dumps(
            {"program": program, "files": list(files), "created": job["created"]}, ensure_ascii=False), "utf-8")
        job.update(status="done", files=list(files), zip_name=f"{pre}_deployment_package.zip",
                   docx_name=f"{pre}_Technical_Specification.docx",
                   outline=[s_["heading"] for k, _ in OUTLINES for s_ in results.get(k, [])],
                   preview={"architecture_png": base64.b64encode(arch_png).decode(),
                            "waveplan_png": base64.b64encode(wave_png).decode() if wave_png else ""})
        step("✅ Package ready")
    except Exception as e:
        logger.exception("ACC package generation failed")
        job.update(status="failed", error=str(e))
        step(f"❌ {e}")


# ════════════════════════════════════════════════════════════════════════════
# Word document
# ════════════════════════════════════════════════════════════════════════════

def _code(doc, text: str):
    from docx.shared import Pt
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    dr.shade(cell, "F3F4F6")
    cell.text = ""
    for i, line in enumerate(str(text).rstrip().splitlines() or [""]):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
    doc.add_paragraph()


def _note(doc, text: str):
    from docx.shared import Pt
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    dr.shade(cell, "FFFDE7")
    cell.text = ""
    dr.add_inline(cell.paragraphs[0], text)
    for r in cell.paragraphs[0].runs:
        r.font.size = Pt(9.5)
    doc.add_paragraph()


def _blocks(doc, blocks: List[Dict[str, Any]]):
    for b in blocks:
        t = b["type"]
        if t == "paragraph":
            dr.add_inline(doc.add_paragraph(), b["text"])
        elif t in ("bullets", "numbered"):
            for it in b["items"]:
                dr.add_inline(doc.add_paragraph(style="List Bullet" if t == "bullets" else "List Number"), it)
        elif t == "table":
            dr.table(doc, b["headers"], b["rows"])
        elif t == "code":
            _code(doc, b["text"])
        elif t == "note":
            _note(doc, b["text"])


def _sections(doc, secs: List[Dict[str, Any]]):
    for s_ in secs:
        dr.heading(doc, s_["heading"], level=s_["level"])
        _blocks(doc, s_["blocks"])


def _figure(doc, png: bytes, caption: str):
    from docx.shared import Inches, Pt
    doc.add_picture(io.BytesIO(png), width=Inches(6.4))
    cap = doc.add_paragraph()
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = dr.GREY


def build_docx(program: Dict[str, Any], secs: Dict[str, List[Dict[str, Any]]], arch_png: bytes,
               wave_png: Optional[bytes], recon: List[Dict[str, str]]) -> bytes:
    from docx.shared import Pt
    doc = dr.new_document()
    pre = program["prefix"]
    total = program.get("total_endpoints")
    tot = f"{total:,}" if total else "[TBD]"
    p = doc.add_paragraph()
    r = p.add_run(f"{program['prefix'].replace('_', ' ')} Technical Specification")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = dr.BRAND
    doc.add_paragraph(f"Date: {date.today().isoformat()}")
    dr.heading(doc, f"Technical Specification: {program['client']} {program['title']}", 1)
    dr.heading(doc, program["product"], 2)

    dr.heading(doc, "1. Document Information", 1)
    dr.table(doc, ["Field", "Detail"], [
        ["Title", f"{program['client']} {program['title']} – Technical Specification"], ["Version", "1.0"],
        ["Date", date.today().isoformat()],
        ["Prepared For", f"{program['client']} – ServiceNow Platform Owners, Endpoint Engineering, Security & "
                         f"Compliance, Regional Deployment Leads" + (f", {program['delivery_partner']} Delivery "
                                                                     f"Leadership" if program['delivery_partner'] else "")],
        ["Deployment Channel", program["deployment_channel"]],
        ["Deployment Timeline", program.get("timeline") or "[TBD]"],
        ["Total Endpoint Target", f"{tot} endpoints (hard reconciliation target)"]])

    _sections(doc, secs.get("intro", []))

    dr.heading(doc, "3. System Architecture & Design", 1)
    dr.heading(doc, "3.1 Architecture Overview", 2)
    dr.add_inline(doc.add_paragraph(), f"Diagram: **{pre}_Architecture.svg** (PlantUML source {pre}_Architecture.puml).")
    try:
        _figure(doc, arch_png, f"Figure 1: {program['architecture']['title']}")
    except Exception:
        _code(doc, arch_diagram.render_ascii(program["architecture"]))
    if program.get("os_split"):
        dr.table(doc, ["Operating system", "Endpoints"],
                 [[o["os"], f"{o['count']:,}" if o.get("count") else "[TBD]"] for o in program["os_split"]])
    _sections(doc, secs.get("arch", []))

    dr.heading(doc, "4. Technical Requirements", 1)
    _sections(doc, secs.get("req", []))
    dr.heading(doc, "4.3 Brand-Tier Endpoint Reconciliation Model", 2)
    _note(doc, f"Rule: the sum of all tiers MUST equal exactly {tot}. Any report, wave count, or dashboard metric "
               "that deviates triggers a reconciliation alert and blocks wave progression.")
    if program["brand_tiers"]:
        known = [b["target"] for b in program["brand_tiers"] if b.get("target")]
        dr.table(doc, ["Brand Tier", "Brands", "Target Devices", "Current Health", "Primary Issues"],
                 [[b["tier"], b["brands"] if b.get("brands") is not None else "—",
                   f"{b['target']:,}" if b.get("target") else "[TBD]", b["health"] or "[TBD]", b["issues"] or "—"]
                  for b in program["brand_tiers"]] +
                 [["TOTAL", sum(b.get("brands") or 0 for b in program["brand_tiers"]), f"{sum(known):,}", "—", "—"]])
    dr.add_inline(doc.add_paragraph(), "**Reconciliation check (computed by the agent):**")
    for rc in recon:
        dr.add_inline(doc.add_paragraph(style="List Bullet"),
                      {"ok": "✔ ", "warning": "⚠ ", "error": "✖ "}.get(rc["level"], "") + rc["message"])
    _code(doc, f"total_deployed = SUM(deployed_count) across all brand tiers\nIF total_deployed != {total or 'TARGET'}:\n"
               f"    raise ReconciliationAlert(expected={total or 'TARGET'}, actual=total_deployed,\n"
               f"                              delta={total or 'TARGET'} - total_deployed, severity=\"HIGH\")\n"
               "    block_wave_progression = true")
    _sections(doc, secs.get("logic", []))
    _sections(doc, secs.get("os", []))

    dr.heading(doc, "6. Deployment Wave Plan", 1)
    dr.add_inline(doc.add_paragraph(), f"Diagram: **{pre}_WavePlan.svg** (PlantUML source {pre}_WavePlan.puml).")
    if wave_png:
        _figure(doc, wave_png, "Figure 2: Wave rollout and stage gates")
    cum = 0
    rows = []
    for w in program["waves"]:
        cum += w.get("target") or 0
        rows.append([w["wave_id"], w["name"], f"{w.get('start') or '[TBD]'} – {w.get('end') or '[TBD]'}",
                     w.get("regions") or "—", f"~{w['target']:,}" if w.get("target") else "[TBD]",
                     f"{cum:,}" if w.get("target") else "—"])
    dr.table(doc, ["Wave", "Name", "Timeline", "Regions", "Target Endpoints", "Cumulative Total"], rows)
    _note(doc, "Wave endpoint counts must be confirmed against live enrollment data before each wave entry gate. "
               f"All waves must reconcile to {tot} total at completion of the last wave.")
    dr.heading(doc, "6.1 Wave Entry & Exit Criteria", 2)
    dr.table(doc, ["Wave", "Entry Criteria", "Exit Criteria"],
             [[w["wave_id"], w.get("entry") or "[TBD]", w.get("exit") or "[TBD]"] for w in program["waves"]])

    dr.heading(doc, "7. Stage Gate Governance", 1)
    _note(doc, "Governance rule: no gate may be set to GO by any automated process. human_approval_required is "
               "permanently true. The system provides a recommendation only; a named human approver must record the "
               "decision with a justification of at least 50 characters. (Enforced by the agent and by a business "
               "rule in ServiceNow.)")
    dr.table(doc, ["Gate", "Name", "Objective", "Owner", "Approver", "Key Criteria", "Unlocks"],
             [[g["gate_id"], g["name"], g.get("objective") or "—", g.get("owner") or "[TBD]",
               g.get("approver") or "[TBD]", g.get("criteria") or "—", g.get("unlocks_wave") or "—"]
              for g in program["gates"]])
    dr.heading(doc, "7.1 Human Approval Audit Log Requirements", 2)
    dr.table(doc, ["Field", "Description"], [
        ["approver_identity", "Full name + corporate email of approver (the signed-in ServiceNow user)"],
        ["approval_timestamp", "UTC datetime of decision"], ["decision_status", "GO / NO-GO / DEFERRED"],
        ["decision_reason", "Free-text justification (minimum 50 characters, mandatory)"],
        ["gate_recommendation", "System-generated recommendation (advisory only)"],
        ["open_risks_acknowledged", "Open risks explicitly acknowledged by the approver"]])

    _sections(doc, secs.get("ops", []))
    _sections(doc, secs.get("ui", []))
    _sections(doc, secs.get("close", []))
    doc.add_paragraph()
    doc.add_paragraph("Generated by SoftProdigy")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ════════════════════════════════════════════════════════════════════════════
# ServiceNow set-up (reuses the Build App engine) + seeding
# ════════════════════════════════════════════════════════════════════════════

BR_GATE = """var p = current.isValidField('decision_status') ? '' : 'u_';
current.setValue(p + 'human_approval_required', true);
var st = String(current.getValue(p + 'decision_status') || 'pending');
if (st != 'pending') {
    var reason = String(current.getValue(p + 'decision_reason') || '');
    if (current.getValue(p + 'approver') == null || current.getValue(p + 'approver') == '' || reason.length < 50) {
        gs.addErrorMessage('Auto-approval blocked. A named human approver and a decision reason of at least 50 characters are required.');
        current.setAbortAction(true);
    }
}"""

BR_AUDIT = """gs.addErrorMessage('Gate decision log entries are immutable and cannot be changed or deleted.');
current.setAbortAction(true);"""

BR_WAVE = """var p = current.isValidField('deployed_endpoints') ? '' : 'u_';
var ga = new GlideAggregate(current.getTableName());
ga.addQuery('sys_id', '!=', current.getUniqueValue());
ga.addAggregate('SUM', p + 'deployed_endpoints');
ga.addAggregate('SUM', p + 'target_endpoints');
ga.query();
var others = 0, targets = 0;
if (ga.next()) {
    others = parseInt(ga.getAggregate('SUM', p + 'deployed_endpoints'), 10) || 0;
    targets = parseInt(ga.getAggregate('SUM', p + 'target_endpoints'), 10) || 0;
}
targets += parseInt(current.getValue(p + 'target_endpoints'), 10) || 0;
var total = others + (parseInt(current.getValue(p + 'deployed_endpoints'), 10) || 0);
if (targets > 0 && total > targets) {
    gs.addErrorMessage('Wave total exceeds the ' + targets + ' endpoint target. Reconcile before proceeding.');
    current.setAbortAction(true);
}"""


def _ch(*vals: str) -> List[Dict[str, str]]:
    return [{"value": v.lower().replace("-", "_").replace(" ", "_"), "label": v} for v in vals]


def setup_plan(app_name: str = "ACC Deployment Program", include_cmdb: bool = True) -> Dict[str, Any]:
    S = lambda n, l, m=100: {"name": n, "label": l, "type": "string", "max_length": m}
    I = lambda n, l: {"name": n, "label": l, "type": "integer"}
    D = lambda n, l: {"name": n, "label": l, "type": "decimal"}
    tables = [
        {"name": "acc_wave", "label": "ACC Wave", "number_prefix": "ACCW", "fields": [
            S("wave_id", "Wave ID", 20), S("wave_name", "Wave name"), {"name": "target_start", "label": "Target start", "type": "date"},
            {"name": "target_end", "label": "Target end", "type": "date"}, S("regions", "Regions", 200),
            I("target_endpoints", "Target endpoints"), I("deployed_endpoints", "Deployed endpoints"),
            D("health_score", "Health score"),
            {"name": "wave_status", "label": "Wave status", "type": "choice", "default_value": "planned",
             "choices": _ch("Planned", "Active", "Complete", "Blocked")},
            S("entry_criteria", "Entry criteria", 4000), S("exit_criteria", "Exit criteria", 4000),
            I("readiness_score", "Readiness score"), S("blockers", "Current blockers", 2000),
            S("executive_notes", "Executive notes", 4000),
            {"name": "owner", "label": "Owner", "type": "reference", "reference": "sys_user"}]},
        {"name": "acc_stage_gate", "label": "ACC Stage Gate", "number_prefix": "ACCG", "fields": [
            S("gate_id", "Gate ID", 20), S("gate_name", "Gate name", 200), S("objective", "Objective", 4000),
            S("owner_name", "Owner", 200), S("required_approver", "Required approver", 200),
            S("key_criteria", "Key criteria", 4000), S("unlocks_wave", "Unlocks wave", 20),
            {"name": "due_date", "label": "Due date", "type": "date"},
            {"name": "decision_status", "label": "Decision status", "type": "choice", "default_value": "pending",
             "choices": _ch("Pending", "GO", "NO-GO", "Deferred")},
            {"name": "approver", "label": "Approver", "type": "reference", "reference": "sys_user"},
            S("approver_identity", "Approver identity", 255),
            {"name": "approval_timestamp", "label": "Approval timestamp", "type": "datetime"},
            S("decision_reason", "Decision reason", 4000), S("gate_recommendation", "Gate recommendation", 4000),
            S("open_risks_acknowledged", "Open risks acknowledged", 4000),
            S("executive_summary", "Executive summary", 4000),
            {"name": "human_approval_required", "label": "Human approval required", "type": "boolean",
             "default_value": "true"}]},
        {"name": "acc_gate_decision", "label": "ACC Gate Decision Log", "fields": [
            {"name": "gate", "label": "Gate", "type": "reference", "reference": "acc_stage_gate"},
            S("gate_id", "Gate ID", 20), S("decision_status", "Decision", 20),
            {"name": "approver", "label": "Approver", "type": "reference", "reference": "sys_user"},
            S("approver_identity", "Approver identity", 255),
            {"name": "decided_at", "label": "Decided at", "type": "datetime"},
            S("decision_reason", "Decision reason", 4000), S("gate_recommendation", "Gate recommendation", 4000),
            S("open_risks_acknowledged", "Open risks acknowledged", 4000)]},
        {"name": "acc_brand_health", "label": "ACC Brand Tier Health", "fields": [
            S("brand_tier", "Brand tier", 40), I("brands", "Brands"), I("target_count", "Target count"),
            I("deployed_count", "Deployed count"), D("reg_success_rate", "Registration success %"),
            D("data_completeness", "Data completeness %"), D("data_freshness", "Data freshness %"),
            {"name": "duplicate_ci_risk", "label": "Duplicate CI risk", "type": "choice", "choices": _ch("Low", "Medium", "High")},
            I("open_remediation", "Open remediation"),
            {"name": "health_status", "label": "Health status", "type": "choice", "choices": _ch("Healthy", "At risk", "Critical")},
            S("current_health", "Current health", 80), S("primary_issues", "Primary issues", 1000),
            D("deploy_confidence", "Deploy confidence")]},
        {"name": "acc_deploy_error", "label": "ACC Deployment Error", "number_prefix": "ACCE", "fields": [
            {"name": "ci", "label": "Configuration item", "type": "reference", "reference": CI_TABLE},
            S("error_code", "Error code", 40), S("message", "Message", 4000), S("region_code", "Region", 20),
            S("wave_id", "Wave ID", 20),
            {"name": "state", "label": "State", "type": "choice", "default_value": "open",
             "choices": _ch("Open", "In progress", "Resolved")},
            {"name": "assigned_to", "label": "Assigned to", "type": "reference", "reference": "sys_user"}]},
    ]
    if include_cmdb:
        tables.append({"name": CI_TABLE, "label": "Computer", "fields": [
            S("region_code", "ACC region", 20), S("wave_id", "ACC wave", 20), S("brand_tier", "Brand tier", 30),
            S("brand_name", "Brand name", 100),
            {"name": "acc_vc_status", "label": "ACC-VC status", "type": "choice",
             "choices": _ch("Registered", "Failed", "Pending", "Stale", "Unregistered")},
            {"name": "intune_enrolled", "label": "Intune enrolled", "type": "boolean"},
            {"name": "duplicate_risk", "label": "Duplicate risk", "type": "choice", "choices": _ch("Low", "Medium", "High")},
            D("deploy_confidence", "Deploy confidence")]})
    return {
        "app_name": app_name, "app_scope": "acc_deploy", "scope_mode": "scoped",
        "description": "ACC-VC endpoint deployment program: waves, human-approved stage gates, audit log, "
                       "brand-tier health and deployment errors (built by SoftProdigy)",
        "roles": [{"name": "acc_program_manager", "description": "Manages ACC deployment waves"},
                  {"name": "acc_gate_approver", "description": "Records stage gate decisions",
                   "contains": ["acc_program_manager"]}],
        "tables": tables,
        "business_rules": [
            {"name": "ACC gate - human approval required", "table": "acc_stage_gate", "when": "before",
             "insert": True, "update": True, "script": BR_GATE,
             "description": "No GO / NO-GO / DEFERRED without a named approver and a 50+ character reason"},
            {"name": "ACC gate decision log is immutable", "table": "acc_gate_decision", "when": "before",
             "insert": False, "update": True, "delete": True, "script": BR_AUDIT},
            {"name": "ACC wave - reconcile deployed endpoints", "table": "acc_wave", "when": "before",
             "insert": True, "update": True, "script": BR_WAVE,
             "description": "Sum of deployed endpoints across waves may not exceed the sum of wave targets"}],
    }


def discover(sn: snc.SNClient) -> Dict[str, Any]:
    q = "^OR".join(f"nameENDSWITH{t}" for t in LOGICAL_TABLES)
    rows = sn.query("sys_db_object", q, "name,sys_id", limit=100)
    tables: Dict[str, Optional[str]] = {}
    for logical in LOGICAL_TABLES:
        c = sorted(r["name"] for r in rows if r.get("name") == f"u_{logical}" or
                   (str(r.get("name", "")).startswith("x_") and str(r["name"]).endswith("_" + logical)))
        tables[logical] = c[0] if c else None
    scope = ""
    if tables.get("acc_wave") and tables["acc_wave"].startswith("x_"):
        scope = tables["acc_wave"][: -len("_acc_wave")]
    fq = "^OR".join(f"elementENDSWITH{f}" for f in CMDB_FIELDS)
    d = sn.query("sys_dictionary", f"name={CI_TABLE}^{fq}", "element", limit=100)
    fields: Dict[str, Optional[str]] = {}
    for f in CMDB_FIELDS:
        els = [r["element"] for r in d if r.get("element") in (f"u_{f}", f"{scope}_{f}" if scope else None) or
               (str(r.get("element", "")).startswith("x_") and str(r["element"]).endswith("_" + f))]
        pref = [e for e in els if scope and e.startswith(scope + "_")]
        fields[f] = (pref or sorted(els) or [None])[0]
    return {"tables": tables, "cmdb_fields": fields, "scope": scope,
            "ready": bool(tables.get("acc_wave") and tables.get("acc_stage_gate")),
            "complete": all(tables.values())}


def fp(table: str) -> str:
    """Field prefix on our own tables: scoped app -> '' ; global -> 'u_'."""
    return "u_" if str(table).startswith("u_") else ""


def seed_program(sn: snc.SNClient, program: Dict[str, Any]) -> Dict[str, Any]:
    m = discover(sn)
    if not m["ready"]:
        raise HTTPException(status_code=409, detail="The ACC tables do not exist yet — run 'Set up in ServiceNow' first")
    out = {"waves": [], "gates": [], "brand_tiers": [], "errors": []}

    def upsert(table: str, key_field: str, key: str, payload: Dict[str, Any], bucket: str):
        p = fp(table)
        data = {p + k: v for k, v in payload.items() if v not in (None, "")}
        try:
            ex = sn.get_one(table, f"{p}{key_field}={key}", "sys_id")
            if ex:
                sn.patch(table, ex["sys_id"], data)
                out[bucket].append({"key": key, "action": "updated", "sys_id": ex["sys_id"], "link": sn.link(table, ex["sys_id"])})
            else:
                rec = sn.create(table, data, fields="sys_id")
                out[bucket].append({"key": key, "action": "created", "sys_id": rec.get("sys_id"),
                                    "link": sn.link(table, rec.get("sys_id", ""))})
        except SNError as e:
            out["errors"].append(f"{bucket[:-1]} {key}: {e}")

    for w in program.get("waves") or []:
        s_, e_ = parse_month(w.get("start")), parse_month(w.get("end"), True)
        upsert(m["tables"]["acc_wave"], "wave_id", w["wave_id"], {
            "wave_id": w["wave_id"], "wave_name": w["name"], "target_start": s_.isoformat() if s_ else None,
            "target_end": e_.isoformat() if e_ else None, "regions": w.get("regions"),
            "target_endpoints": w.get("target"), "entry_criteria": w.get("entry"), "exit_criteria": w.get("exit")},
            "waves")
    waves = {w["wave_id"]: w for w in program.get("waves") or []}
    for g in program.get("gates") or []:
        due = parse_month((waves.get(g.get("unlocks_wave")) or {}).get("start"))
        existing = sn.get_one(m["tables"]["acc_stage_gate"], f"{fp(m['tables']['acc_stage_gate'])}gate_id={g['gate_id']}", "sys_id")
        payload = {"gate_id": g["gate_id"], "gate_name": g["name"], "objective": g.get("objective"),
                   "owner_name": g.get("owner"), "required_approver": g.get("approver"),
                   "key_criteria": g.get("criteria"), "unlocks_wave": g.get("unlocks_wave"),
                   "due_date": due.isoformat() if due else None, "human_approval_required": "true"}
        if not existing:   # never reset a decision that a human already recorded
            payload["decision_status"] = "pending"
        upsert(m["tables"]["acc_stage_gate"], "gate_id", g["gate_id"], payload, "gates")
    if m["tables"].get("acc_brand_health"):
        for b in program.get("brand_tiers") or []:
            upsert(m["tables"]["acc_brand_health"], "brand_tier", b["tier"], {
                "brand_tier": b["tier"], "brands": b.get("brands"), "target_count": b.get("target"),
                "current_health": b.get("health"), "primary_issues": b.get("issues")}, "brand_tiers")
    out["tables"] = m["tables"]
    return out


# ════════════════════════════════════════════════════════════════════════════
# Command Center (live, read-only)
# ════════════════════════════════════════════════════════════════════════════

def _pct(a: Optional[int], b: Optional[int]) -> Optional[float]:
    if a is None or not b:
        return None
    return round(100.0 * a / b, 1)


def _status(value: Optional[float], good: float, warn: float, higher_is_better: bool = True) -> str:
    if value is None:
        return "unknown"
    if higher_is_better:
        return "healthy" if value >= good else "at_risk" if value >= warn else "critical"
    return "healthy" if value <= good else "at_risk" if value <= warn else "critical"


def _rows(sn, table, query, fields, limit=100, order=""):
    q = "^".join(x for x in (query, f"ORDERBY{order}" if order else "") if x)
    return sn.query(table, q, fields, limit=limit)


def command_center(sn: snc.SNClient, target_override: Optional[int] = None) -> Dict[str, Any]:
    notes: List[str] = []
    alerts: List[Dict[str, str]] = []
    m = discover(sn)
    F = m["cmdb_fields"]
    T = m["tables"]
    today = date.today()

    def safe(label, fn, default=None):
        try:
            return fn()
        except SNError as e:
            notes.append(f"{label}: could not read ({e.status} {e.message})")
            return default

    base = "install_status=1"
    total_ci = safe("CMDB computers", lambda: sn.count(CI_TABLE, base), 0) or 0
    all_ci = safe("CMDB computers", lambda: sn.count(CI_TABLE, ""), 0) or 0
    fresh24 = safe("freshness", lambda: sn.count(CI_TABLE, base + "^last_discovered>=javascript:gs.hoursAgoStart(24)"))
    stale48 = safe("freshness", lambda: sn.count(CI_TABLE, base + "^last_discovered<javascript:gs.hoursAgoStart(48)"
                                                         "^ORlast_discoveredISEMPTY"))
    complete = safe("completeness", lambda: sn.count(CI_TABLE, base + "^serial_numberISNOTEMPTY^osISNOTEMPTY^nameISNOTEMPTY"))
    os_split = safe("OS split", lambda: sorted(
        [{"os": v.get("os") or "(empty)", "count": n} for v, n in sn.group_count(CI_TABLE, "os", base)],
        key=lambda x: -x["count"])[:12], [])
    dup_serial = safe("duplicate serials", lambda: sn.group_count(
        CI_TABLE, "serial_number", base + "^serial_numberISNOTEMPTY", having="count^serial_number^>^1"), []) or []
    dup_name = safe("duplicate names", lambda: sn.group_count(
        CI_TABLE, "name", base + "^nameISNOTEMPTY", having="count^name^>^1"), []) or []
    dup_serial = [(v, n) for v, n in dup_serial if n > 1]
    dup_name = [(v, n) for v, n in dup_name if n > 1]
    dup_records = sum(n - 1 for _, n in dup_serial)

    acc_plugin = safe("ACC plugin", lambda: sn.table_exists(ACC_AGENT_TABLE), False)
    agents = agents_24 = None
    if acc_plugin:
        agents = safe("ACC agents", lambda: sn.count(ACC_AGENT_TABLE, ""))
        agents_24 = safe("ACC agents", lambda: sn.count(ACC_AGENT_TABLE, "sys_updated_on>=javascript:gs.hoursAgoStart(24)"))

    status_split: Dict[str, int] = {}
    if F.get("acc_vc_status"):
        for v, n in safe("ACC status", lambda: sn.group_count(CI_TABLE, F["acc_vc_status"], base), []) or []:
            status_split[(v.get(F["acc_vc_status"]) or "(not set)").lower()] = n
    registered = status_split.get("registered")
    failed = status_split.get("failed")
    tracked = sum(n for k, n in status_split.items() if k not in ("(not set)", "unregistered")) if status_split else None

    if tracked:
        deployed, deployed_src = tracked, f"CIs with an ACC-VC status ({F['acc_vc_status']})"
    elif agents is not None:
        deployed, deployed_src = agents, f"ACC agent records ({ACC_AGENT_TABLE})"
    else:
        deployed, deployed_src = total_ci, "active computers in cmdb_ci_computer (no ACC tracking data yet)"
        notes.append("No ACC-VC status field or ACC plugin data found — coverage uses all active computers.")

    # waves / gates / brand tiers
    waves, gates, brands = [], [], []
    if T.get("acc_wave"):
        p = fp(T["acc_wave"])
        for r in safe("waves", lambda: _rows(sn, T["acc_wave"], "", "", 50, p + "wave_id"), []) or []:
            waves.append({"sys_id": r.get("sys_id"), "wave_id": r.get(p + "wave_id"), "name": r.get(p + "wave_name"),
                          "start": r.get(p + "target_start"), "end": r.get(p + "target_end"),
                          "regions": r.get(p + "regions"), "target": _int(r.get(p + "target_endpoints")),
                          "deployed": _int(r.get(p + "deployed_endpoints")) or 0,
                          "status": r.get(p + "wave_status") or "planned", "health": r.get(p + "health_score"),
                          "entry": r.get(p + "entry_criteria"), "exit": r.get(p + "exit_criteria"),
                          "blockers": r.get(p + "blockers"), "notes": r.get(p + "executive_notes"),
                          "readiness": r.get(p + "readiness_score"), "link": sn.link(T["acc_wave"], r.get("sys_id", ""))})
    if T.get("acc_stage_gate"):
        p = fp(T["acc_stage_gate"])
        for r in safe("gates", lambda: _rows(sn, T["acc_stage_gate"], "", "", 50, p + "gate_id"), []) or []:
            gates.append({"sys_id": r.get("sys_id"), "gate_id": r.get(p + "gate_id"), "name": r.get(p + "gate_name"),
                          "objective": r.get(p + "objective"), "owner": r.get(p + "owner_name"),
                          "approver_required": r.get(p + "required_approver"), "criteria": r.get(p + "key_criteria"),
                          "unlocks_wave": r.get(p + "unlocks_wave"), "due": r.get(p + "due_date"),
                          "status": r.get(p + "decision_status") or "pending",
                          "approver_identity": r.get(p + "approver_identity"),
                          "decided_at": r.get(p + "approval_timestamp"), "reason": r.get(p + "decision_reason"),
                          "risks": r.get(p + "open_risks_acknowledged"),
                          "recommendation": r.get(p + "gate_recommendation"),
                          "link": sn.link(T["acc_stage_gate"], r.get("sys_id", ""))})
    if T.get("acc_brand_health"):
        p = fp(T["acc_brand_health"])
        for r in safe("brand tiers", lambda: _rows(sn, T["acc_brand_health"], "", "", 50), []) or []:
            brands.append({"tier": r.get(p + "brand_tier"), "brands": r.get(p + "brands"),
                           "target": _int(r.get(p + "target_count")), "health": r.get(p + "current_health"),
                           "issues": r.get(p + "primary_issues")})
    live_tiers = {}
    if F.get("brand_tier"):
        for v, n in safe("brand tiers", lambda: sn.group_count(CI_TABLE, F["brand_tier"], base), []) or []:
            live_tiers[(v.get(F["brand_tier"]) or "(not set)").upper()] = n
    for b in brands:
        b["deployed_live"] = live_tiers.get(str(b["tier"] or "").upper())

    regions = []
    if F.get("region_code"):
        rc = F["region_code"]
        counts = {(v.get(rc) or "(not set)"): n for v, n in safe("regions", lambda: sn.group_count(CI_TABLE, rc, base), []) or []}
        fresh_r = {(v.get(rc) or "(not set)"): n for v, n in safe("regions", lambda: sn.group_count(
            CI_TABLE, rc, base + "^last_discovered>=javascript:gs.hoursAgoStart(24)"), []) or []}
        reg_r: Dict[str, Dict[str, int]] = {}
        if F.get("acc_vc_status"):
            for v, n in safe("regions", lambda: sn.group_count(CI_TABLE, f"{rc},{F['acc_vc_status']}", base), []) or []:
                reg_r.setdefault(v.get(rc) or "(not set)", {})[(v.get(F["acc_vc_status"]) or "").lower()] = n
        for reg, n in sorted(counts.items(), key=lambda x: -x[1]):
            rr = reg_r.get(reg, {})
            fr = _pct(fresh_r.get(reg, 0), n)
            succ = _pct(rr.get("registered"), (rr.get("registered") or 0) + (rr.get("failed") or 0)) if rr else None
            worst = min([x for x in (fr, succ) if x is not None] or [None]) if (fr is not None or succ is not None) else None
            regions.append({"region": reg, "deployed": n, "registered": rr.get("registered"), "failed": rr.get("failed"),
                            "fresh_pct": fr, "success_pct": succ,
                            "health": _status(worst, 90, 75) if worst is not None else "unknown"})

    open_errors = None
    if T.get("acc_deploy_error"):
        open_errors = safe("deployment errors", lambda: sn.count(T["acc_deploy_error"], f"{fp(T['acc_deploy_error'])}state!=resolved"))

    # target
    wave_target = sum(w["target"] or 0 for w in waves) or None
    target = target_override or wave_target
    target_src = "entered on screen" if target_override else ("sum of wave targets" if wave_target else "not set")

    expected_pct = None
    if waves and target:
        done_target = 0
        for w in waves:
            try:
                end = date.fromisoformat(str(w["end"])[:10]) if w.get("end") else None
            except ValueError:
                end = None
            if end and end < today:
                done_target += w["target"] or 0
        expected_pct = _pct(done_target, target)

    coverage = _pct(deployed, target) if target else None
    kpis = [
        {"key": "coverage", "label": "Deployment Coverage", "value": coverage, "unit": "%",
         "detail": f"{deployed:,} of {target:,}" if target else f"{deployed:,} deployed (no target set)",
         "threshold": f"≥ {expected_pct}% (waves due by today)" if expected_pct is not None else "≥ wave target %",
         "status": (_status(coverage, expected_pct, expected_pct * 0.8) if coverage is not None and expected_pct
                    else "info" if coverage is not None else "unknown"), "source": deployed_src},
        {"key": "reg_success", "label": "Registration Success Rate",
         "value": _pct(registered, (registered or 0) + (failed or 0)) if (registered or failed) else None, "unit": "%",
         "detail": f"{registered or 0:,} registered · {failed or 0:,} failed" if status_split else "no ACC-VC status data",
         "threshold": "≥ 90%", "source": "ACC-VC status on cmdb_ci_computer"},
        {"key": "completeness", "label": "Data Completeness", "value": _pct(complete, total_ci), "unit": "%",
         "detail": f"{complete or 0:,} of {total_ci:,} have name + serial + OS", "threshold": "≥ 85%",
         "source": "cmdb_ci_computer"},
        {"key": "freshness", "label": "Data Freshness", "value": _pct(fresh24, total_ci), "unit": "%",
         "detail": f"{fresh24 or 0:,} discovered in the last 24h · {stale48 or 0:,} stale > 48h", "threshold": "≥ 90%",
         "source": "last_discovered"},
        {"key": "duplicates", "label": "Duplicate CI %", "value": _pct(dup_records, total_ci), "unit": "%",
         "detail": f"{len(dup_serial)} duplicated serial numbers ({dup_records} extra CIs) · {len(dup_name)} duplicated names",
         "threshold": "≤ 2%", "source": "serial_number groups"},
        {"key": "comm_health", "label": "Endpoint Communication Health", "value": _pct(agents_24, agents), "unit": "%",
         "detail": f"{agents_24 or 0:,} of {agents:,} ACC agent records updated in 24h" if agents else
         "ACC plugin not installed / no agents", "threshold": "≥ 95%", "source": ACC_AGENT_TABLE},
        {"key": "waves", "label": "Wave Completion", "value": _pct(sum(1 for w in waves if w["status"] == "complete"), len(waves)),
         "unit": "%", "detail": f"{sum(1 for w in waves if w['status'] == 'complete')} of {len(waves)} waves complete"
         if waves else "program not set up", "threshold": "per schedule", "source": T.get("acc_wave") or "—"},
        {"key": "gates", "label": "Stage Gates Approved", "value": _pct(sum(1 for g in gates if g["status"] == "go"), len(gates)),
         "unit": "%", "detail": f"{sum(1 for g in gates if g['status'] == 'go')} of {len(gates)} GO (human-approved)"
         if gates else "program not set up", "threshold": "per schedule", "source": T.get("acc_stage_gate") or "—"},
    ]
    if F.get("intune_enrolled"):
        enrolled = safe("Intune", lambda: sn.count(CI_TABLE, f"{base}^{F['intune_enrolled']}=true"))
        kpis.append({"key": "intune", "label": "Intune Enrolled", "value": _pct(enrolled, total_ci), "unit": "%",
                     "detail": f"{enrolled or 0:,} of {total_ci:,}", "threshold": "100% of wave scope",
                     "source": F["intune_enrolled"]})
    for k in kpis:
        if "status" not in k:
            th = {"reg_success": (90, 75, True), "completeness": (85, 70, True), "freshness": (90, 70, True),
                  "duplicates": (2, 5, False), "comm_health": (95, 85, True), "intune": (95, 80, True)}.get(k["key"])
            k["status"] = _status(k["value"], th[0], th[1], th[2]) if th else ("info" if k["value"] is not None else "unknown")

    # alerts
    if target and wave_target and target != wave_target:
        alerts.append({"level": "error", "message": f"Reconciliation: waves add up to {wave_target:,}, target is {target:,} "
                                                    f"(delta {target - wave_target:+,})."})
    deployed_waves = sum(w["deployed"] for w in waves)
    if target and deployed_waves > target:
        alerts.append({"level": "error", "message": f"Wave deployed counts ({deployed_waves:,}) exceed the {target:,} target."})
    if target and deployed > target:
        alerts.append({"level": "warning", "message": f"{deployed:,} deployed endpoints is more than the {target:,} target — "
                                                      "possible duplicates or out-of-scope devices."})
    if dup_serial:
        alerts.append({"level": "warning", "message": f"{len(dup_serial)} serial numbers are shared by more than one active CI."})
    if total_ci and stale48 is not None and stale48 / total_ci > 0.3:
        alerts.append({"level": "warning", "message": f"{_pct(stale48, total_ci)}% of active computers not seen for 48h+ "
                                                      "(rollback trigger threshold: 30%)."})
    for g in gates:
        try:
            due = date.fromisoformat(str(g["due"])[:10]) if g.get("due") else None
        except ValueError:
            due = None
        if g["status"] == "pending" and due and due < today:
            alerts.append({"level": "warning", "message": f"{g['gate_id']} decision overdue since {due.isoformat()}."})
    go = {g["gate_id"] for g in gates if g["status"] == "go"}
    order = [w["wave_id"] for w in waves]
    for w in waves:
        if w["status"] == "active":
            need = [g["gate_id"] for g in gates if g.get("unlocks_wave") in order and
                    order.index(g["unlocks_wave"]) <= order.index(w["wave_id"]) and g["gate_id"] not in go]
            if need:
                alerts.append({"level": "error", "message": f"{w['wave_id']} is active but {', '.join(need)} not approved."})
        if w["status"] == "blocked":
            alerts.append({"level": "warning", "message": f"{w['wave_id']} is blocked: {w.get('blockers') or 'no reason given'}"})

    return {"generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"), "instance": sn.instance,
            "setup": m, "acc_plugin": bool(acc_plugin), "target": target, "target_source": target_src,
            "totals": {"active_computers": total_ci, "all_computers": all_ci, "deployed": deployed,
                       "deployed_source": deployed_src, "fresh24": fresh24, "stale48": stale48, "agents": agents,
                       "open_errors": open_errors},
            "kpis": kpis, "os_split": os_split, "status_split": status_split, "regions": regions,
            "waves": waves, "gates": gates, "brand_tiers": brands, "live_brand_tiers": live_tiers,
            "duplicates": {"serial_groups": [{"serial_number": v.get("serial_number"), "count": n} for v, n in dup_serial[:50]],
                           "name_groups": [{"name": v.get("name"), "count": n} for v, n in dup_name[:50]],
                           "extra_records": dup_records},
            "alerts": alerts, "notes": notes}


# ════════════════════════════════════════════════════════════════════════════
# Routes
# ════════════════════════════════════════════════════════════════════════════

def _client(body: dict) -> snc.SNClient:
    return snc.client_from(body)


async def _run(fn, *a, what="ServiceNow call"):
    try:
        return await asyncio.to_thread(fn, *a)
    except SNError as e:
        raise e.to_http(what)


@acc_router.get("/snow/acc/default-program")
async def default_program():
    p = normalize_program(json.loads(json.dumps(DEFAULT_PROGRAM)))
    return {"program": p, "reconciliation": reconcile(p)}


@acc_router.post("/snow/acc/generate")
async def acc_generate(file: Optional[UploadFile] = File(None), text: Optional[str] = Form(None),
                       client: Optional[str] = Form(None), delivery_partner: Optional[str] = Form(None),
                       total_endpoints: Optional[str] = Form(None), use_template: Optional[str] = Form(None)):
    """Start generating the deployment package. Poll /snow/acc/generate/{job_id}."""
    from snow_universal import file_to_text
    source = (text or "").strip()
    if file is not None and file.filename:
        got = await file_to_text(file)
        source = (source + "\n\n" + (got.get("text") or "")).strip()
    if not source and str(use_template).lower() != "true":
        raise HTTPException(status_code=400, detail="Upload the functional spec / notes, paste text, or choose the template")
    if len(source) > MAX_SOURCE:
        source = source[:MAX_SOURCE]
    overrides = {"client": (client or "").strip(), "delivery_partner": (delivery_partner or "").strip(),
                 "total_endpoints": _int(total_endpoints) if total_endpoints else None}
    job_id = uuid.uuid4().hex[:10]
    job = {"job_id": job_id, "status": "queued", "progress": [], "created": datetime.now().isoformat(timespec="seconds"),
           "source_chars": len(source)}
    with _gen_lock:
        _gen_jobs[job_id] = job
    threading.Thread(target=generate_package, args=(job, source, overrides), daemon=True).start()
    return {"job_id": job_id, "status": "queued"}


@acc_router.get("/snow/acc/generate/{job_id}")
async def acc_generate_status(job_id: str):
    job = _gen_jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Unknown package job (the agent may have restarted)")
    return job


def _pkg_path(job_id: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{10}", job_id or ""):
        raise HTTPException(status_code=400, detail="bad package id")
    p = PKG_DIR / f"{job_id}.zip"
    if not p.exists():
        raise HTTPException(status_code=404, detail="Package not found — generate it again")
    return p


@acc_router.get("/snow/acc/package/{job_id}")
async def acc_package(job_id: str, file: Optional[str] = None):
    """The whole package as .zip, or one file of it (?file=<name>)."""
    p = _pkg_path(job_id)
    with zipfile.ZipFile(p) as z:
        names = z.namelist()
        if file:
            if file not in names:
                raise HTTPException(status_code=404, detail="file not in package")
            data = z.read(file)
            media = {"docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                     "png": "image/png", "svg": "image/svg+xml", "json": "application/json"}.get(file.rsplit(".", 1)[-1],
                                                                                              "text/plain")
            return Response(content=data, media_type=media,
                            headers={"Content-Disposition": f'attachment; filename="{file}"'})
    pre = next((n[: -len("_Technical_Specification.docx")] for n in names if n.endswith("_Technical_Specification.docx")), "ACC")
    return Response(content=p.read_bytes(), media_type="application/zip",
                    headers={"Content-Disposition": f'attachment; filename="{pre}_deployment_package.zip"'})


@acc_router.post("/snow/acc/status")
async def acc_status(body: dict):
    sn = _client(body)
    return await _run(discover, sn, what="ServiceNow (ACC set-up check)")


@acc_router.post("/snow/acc/setup")
async def acc_setup(body: dict):
    """Create the ACC program model in ServiceNow through the Build App engine (scope + update set)."""
    sn = _client(body)
    try:
        await asyncio.to_thread(sn.user)
    except SNError as e:
        raise e.to_http("ServiceNow login")
    plan = setup_plan(body.get("app_name") or "ACC Deployment Program", body.get("include_cmdb_fields", True) is not False)
    for k in ("scope_mode", "existing_app"):
        if body.get(k):
            plan[k] = body[k]
    opts = {"update_set_name": body.get("update_set_name") or f"HAI_ACC_Deployment_{datetime.now():%Y%m%d_%H%M}",
            "update_set_description": "ACC-VC deployment program (waves, stage gates, audit log) — SoftProdigy",
            "stop_on_error": False, "rollback_on_error": False,
            "restore_previous_context": bool(body.get("restore_previous_context"))}
    if body.get("update_set_sys_id"):
        opts["update_set_sys_id"] = body["update_set_sys_id"]
    job_id = sn_build.start_build(sn, plan, opts)
    threading.Thread(target=sn_build.run_build, args=(job_id, sn, plan, opts), daemon=True).start()
    return {"job_id": job_id, "status": "queued", "plan_counts": {k: len(plan.get(k) or []) for k in
                                                                  ("roles", "tables", "business_rules")}}


@acc_router.post("/snow/acc/seed")
async def acc_seed(body: dict):
    sn = _client(body)
    program = normalize_program(body.get("program") or json.loads(json.dumps(DEFAULT_PROGRAM)))
    if not body.get("program") and body.get("package_id"):
        meta = PKG_DIR / f"{body['package_id']}.json"
        if meta.exists():
            program = normalize_program(json.loads(meta.read_text("utf-8")).get("program"))
    return await _run(seed_program, sn, program, what="ServiceNow (seed program)")


@acc_router.post("/snow/acc/command-center")
async def acc_command_center(body: dict):
    sn = _client(body)
    return await _run(command_center, sn, _int(body.get("target_endpoints")), what="ServiceNow (command center)")


def _brief_input(data: Dict[str, Any]) -> str:
    keep = {k: data.get(k) for k in ("target", "totals", "kpis", "os_split", "status_split", "regions", "waves",
                                     "gates", "brand_tiers", "alerts")}
    for w in keep.get("waves") or []:
        w.pop("link", None)
    for g in keep.get("gates") or []:
        g.pop("link", None)
    keep["duplicates"] = {k: (data.get("duplicates") or {}).get(k) for k in ("extra_records",)}
    keep["duplicates"]["top_serials"] = ((data.get("duplicates") or {}).get("serial_groups") or [])[:10]
    return json.dumps(keep, ensure_ascii=False, default=str)[:24000]


@acc_router.post("/snow/acc/briefing")
async def acc_briefing(body: dict):
    """AI Remediation Briefing from live Command Center data (advisory only)."""
    data = body.get("data")
    if not isinstance(data, dict):
        sn = _client(body)
        data = await _run(command_center, sn, _int(body.get("target_endpoints")), what="ServiceNow (command center)")
    prompt = ("ACC_BRIEFING\nLive ACC-VC deployment data from ServiceNow:\n" + _brief_input(data) + """

Write the AI Remediation Briefing as JSON:
{"confidence": 0-100 overall deployment confidence,
 "risk_trend": "improving|stable|worsening|unknown",
 "key_risks": [{"risk": "...", "severity": "high|medium|low", "evidence": "number from the data"}],   (max 5, ranked)
 "regions_behind": ["..."], "duplicate_patterns": ["..."], "stale_data_issues": ["..."],
 "actions": [{"action": "specific step", "owner": "resolver group / role", "due": "e.g. 24h"}],
 "next_48h": ["prioritised action"],
 "human_approval_items": ["decision a named human must make"],
 "go_no_go": "advisory recommendation for the next pending gate, stating it requires human approval",
 "executive_summary": "4-6 sentences for leadership"}
Only use facts present in the data; say 'no data' when a metric is missing. Never state that a gate is approved.""")
    try:
        out = await groq_client.agenerate_json(prompt, system=SYSTEM, max_tokens=3500)
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"AI service unavailable: {e}")
    if not isinstance(out, dict):
        raise HTTPException(status_code=502, detail="AI returned an unreadable briefing, please retry")
    out["advisory_only"] = True
    return out


def _gate_list(sn, m) -> List[dict]:
    t = m["tables"]["acc_stage_gate"]
    p = fp(t)
    return [{"sys_id": r["sys_id"], "gate_id": r.get(p + "gate_id"), "status": r.get(p + "decision_status") or "pending",
             "unlocks_wave": r.get(p + "unlocks_wave"), "name": r.get(p + "gate_name"),
             "criteria": r.get(p + "key_criteria"), "objective": r.get(p + "objective"),
             "required_approver": r.get(p + "required_approver")}
            for r in sn.query(t, f"ORDERBY{p}gate_id", "", limit=50)]


@acc_router.post("/snow/acc/gate/recommend")
async def acc_gate_recommend(body: dict):
    """Advisory GO / NO-GO for one gate, checked against live data. Never records a decision."""
    sn = _client(body)
    m = await _run(discover, sn)
    if not m["ready"]:
        raise HTTPException(status_code=409, detail="Set up the ACC program in ServiceNow first")
    gates = await _run(_gate_list, sn, m)
    gate = next((g for g in gates if g["sys_id"] == body.get("gate_sys_id")), None)
    if not gate:
        raise HTTPException(status_code=404, detail="Gate not found")
    data = body.get("data") if isinstance(body.get("data"), dict) else await _run(command_center, sn, None)
    prompt = ("ACC_GATE_RECOMMEND\nGate under review:\n" + json.dumps(gate, ensure_ascii=False) +
              "\n\nLive deployment data:\n" + _brief_input(data) + """

Check each key criterion of the gate against the data. JSON:
{"recommendation": "GO|NO-GO|DEFER", "confidence": 0-100,
 "criteria": [{"criterion": "...", "status": "met|not_met|unknown", "evidence": "..."}],
 "open_risks": ["..."], "summary": "2-3 sentences"}
Missing data = 'unknown', never 'met'. This is advisory; a named human must decide.""")
    try:
        out = await groq_client.agenerate_json(prompt, system=SYSTEM, max_tokens=2500)
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"AI service unavailable: {e}")
    if not isinstance(out, dict):
        raise HTTPException(status_code=502, detail="AI returned an unreadable answer, please retry")
    out["advisory_only"] = True
    out["gate_id"] = gate["gate_id"]
    return out


def _decide(sn: snc.SNClient, body: dict) -> dict:
    m = discover(sn)
    if not m["ready"]:
        raise HTTPException(status_code=409, detail="Set up the ACC program in ServiceNow first")
    decision = str(body.get("decision") or "").lower().replace("-", "_").replace(" ", "_")
    if decision not in ("go", "no_go", "deferred"):
        raise HTTPException(status_code=400, detail="decision must be GO, NO-GO or DEFERRED")
    reason = str(body.get("reason") or "").strip()
    if len(reason) < 50:
        raise HTTPException(status_code=400, detail="Decision reason is required. Provide a justification of at least "
                                                    f"50 characters (currently {len(reason)}).")
    if body.get("confirm") is not True:
        raise HTTPException(status_code=400, detail="Human confirmation required (confirm=true)")
    gates = _gate_list(sn, m)
    idx = next((i for i, g in enumerate(gates) if g["sys_id"] == body.get("gate_sys_id")), None)
    if idx is None:
        raise HTTPException(status_code=404, detail="Gate not found")
    gate = gates[idx]
    if decision == "go":
        not_go = [g["gate_id"] for g in gates[:idx] if g["status"] != "go"]
        if not_go:
            raise HTTPException(status_code=409, detail=f"{gate['gate_id']} cannot be GO while earlier gate(s) "
                                                        f"{', '.join(not_go)} are not GO.")
    u = sn.user()
    identity = f"{u.get('name') or u.get('user_name')} <{u.get('email') or 'no email'}> ({u.get('user_name')})"
    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    t = m["tables"]["acc_stage_gate"]
    p = fp(t)
    payload = {p + "decision_status": decision, p + "approver": u["sys_id"], p + "approver_identity": identity,
               p + "approval_timestamp": ts, p + "decision_reason": reason,
               p + "open_risks_acknowledged": str(body.get("open_risks_acknowledged") or "")[:4000],
               p + "gate_recommendation": str(body.get("recommendation") or "")[:4000],
               p + "human_approval_required": "true"}
    sn.patch(t, gate["sys_id"], payload)
    after = sn.get(t, gate["sys_id"], f"{p}decision_status,{p}approver_identity")
    if (after.get(p + "decision_status") or "") != decision:
        raise HTTPException(status_code=409, detail="ServiceNow did not accept the decision (a business rule blocked it). "
                                                    "Check the gate record for the error message.")
    audit = None
    if m["tables"].get("acc_gate_decision"):
        a = m["tables"]["acc_gate_decision"]
        q = fp(a)
        audit = sn.create(a, {q + "gate": gate["sys_id"], q + "gate_id": gate["gate_id"], q + "decision_status": decision,
                              q + "approver": u["sys_id"], q + "approver_identity": identity, q + "decided_at": ts,
                              q + "decision_reason": reason,
                              q + "open_risks_acknowledged": str(body.get("open_risks_acknowledged") or "")[:4000],
                              q + "gate_recommendation": str(body.get("recommendation") or "")[:4000]}, fields="sys_id")
    return {"gate_id": gate["gate_id"], "decision": decision, "approver_identity": identity, "timestamp_utc": ts,
            "gate_link": sn.link(t, gate["sys_id"]),
            "audit_link": sn.link(m["tables"]["acc_gate_decision"], audit["sys_id"]) if audit else None}


@acc_router.post("/snow/acc/gate/decide")
async def acc_gate_decide(body: dict):
    """Record a human gate decision (GO / NO-GO / DEFERRED) with approver identity, reason and audit row."""
    sn = _client(body)
    return await _run(_decide, sn, body, what="ServiceNow (gate decision)")


def _audit(sn: snc.SNClient, gate_sys_id: str = "") -> List[dict]:
    m = discover(sn)
    a = m["tables"].get("acc_gate_decision")
    if not a:
        return []
    q = fp(a)
    rows = sn.query(a, (f"{q}gate={gate_sys_id}^" if gate_sys_id else "") + f"ORDERBYDESC{q}decided_at", "", limit=100)
    return [{"gate_id": r.get(q + "gate_id"), "decision": r.get(q + "decision_status"),
             "approver_identity": r.get(q + "approver_identity"), "decided_at": r.get(q + "decided_at"),
             "reason": r.get(q + "decision_reason"), "risks": r.get(q + "open_risks_acknowledged"),
             "link": sn.link(a, r.get("sys_id", ""))} for r in rows]


@acc_router.post("/snow/acc/gate/audit")
async def acc_gate_audit(body: dict):
    sn = _client(body)
    return {"entries": await _run(_audit, sn, body.get("gate_sys_id") or "")}


def _wave_update(sn: snc.SNClient, body: dict) -> dict:
    m = discover(sn)
    if not m["ready"]:
        raise HTTPException(status_code=409, detail="Set up the ACC program in ServiceNow first")
    t = m["tables"]["acc_wave"]
    p = fp(t)
    waves = sn.query(t, f"ORDERBY{p}wave_id", "", limit=50)
    idx = next((i for i, w in enumerate(waves) if w["sys_id"] == body.get("wave_sys_id")), None)
    if idx is None:
        raise HTTPException(status_code=404, detail="Wave not found")
    w = waves[idx]
    payload: Dict[str, Any] = {}
    status = body.get("status")
    if status:
        status = str(status).lower()
        if status not in ("planned", "active", "complete", "blocked"):
            raise HTTPException(status_code=400, detail="status must be planned, active, complete or blocked")
        if status == "active":
            order = [x.get(p + "wave_id") for x in waves]
            gates = _gate_list(sn, m)
            need = [g["gate_id"] for g in gates if g.get("unlocks_wave") in order and
                    order.index(g["unlocks_wave"]) <= idx and g["status"] != "go"]
            if need:
                raise HTTPException(status_code=409, detail=f"{w.get(p + 'wave_id')} cannot start: stage gate(s) "
                                                            f"{', '.join(need)} are not approved (GO) by a human.")
        if status == "blocked" and not (body.get("blockers") or w.get(p + "blockers")):
            raise HTTPException(status_code=400, detail="Say what is blocking the wave")
        payload[p + "wave_status"] = status
    if body.get("deployed_endpoints") not in (None, ""):
        dep = _int(body["deployed_endpoints"])
        if dep is None or dep < 0:
            raise HTTPException(status_code=400, detail="deployed_endpoints must be a whole number ≥ 0")
        others = sum(_int(x.get(p + "deployed_endpoints")) or 0 for x in waves if x["sys_id"] != w["sys_id"])
        target = _int(body.get("target_endpoints")) or sum(_int(x.get(p + "target_endpoints")) or 0 for x in waves)
        if target and others + dep > target:
            raise HTTPException(status_code=409, detail=f"Wave total exceeds the {target:,} endpoint target "
                                                        f"({others + dep:,}). Reconcile before proceeding.")
        payload[p + "deployed_endpoints"] = dep
    for k, f in (("blockers", "blockers"), ("executive_notes", "executive_notes"), ("health_score", "health_score"),
                 ("readiness_score", "readiness_score")):
        if body.get(k) not in (None, ""):
            payload[p + f] = str(body[k])[:4000]
    if not payload:
        raise HTTPException(status_code=400, detail="Nothing to update")
    if body.get("confirm") is not True:
        raise HTTPException(status_code=400, detail="Confirmation required (confirm=true)")
    sn.patch(t, w["sys_id"], payload)
    return {"wave_id": w.get(p + "wave_id"), "updated": list(payload), "link": sn.link(t, w["sys_id"])}


@acc_router.post("/snow/acc/wave/update")
async def acc_wave_update(body: dict):
    sn = _client(body)
    return await _run(_wave_update, sn, body, what="ServiceNow (wave update)")


def _duplicates(sn: snc.SNClient, limit: int = 20) -> List[dict]:
    groups = [(v, n) for v, n in sn.group_count(CI_TABLE, "serial_number", "install_status=1^serial_numberISNOTEMPTY",
                                                 having="count^serial_number^>^1") if n > 1][:limit]
    out = []
    for v, n in groups:
        serial = v.get("serial_number")
        cis = sn.query(CI_TABLE, f"install_status=1^serial_number={serial}", "sys_id,name,os,last_discovered,sys_created_on", 10)
        out.append({"serial_number": serial, "count": n,
                    "cis": [dict(c, link=sn.link(CI_TABLE, c.get("sys_id", ""))) for c in cis]})
    return out


@acc_router.post("/snow/acc/duplicates")
async def acc_duplicates(body: dict):
    sn = _client(body)
    return {"groups": await _run(_duplicates, sn, what="ServiceNow (duplicate CIs)")}


@acc_router.post("/snow/acc/report")
async def acc_report(body: dict):
    """Excel export of the Command Center view the user is looking at (+ briefing when present)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    data = body.get("data") or {}
    wb = Workbook()
    head = PatternFill("solid", fgColor="1E5EBF")

    def sheet(title, headers, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title[:31]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head
        for r in rows:
            ws.append([("" if v is None else v) for v in r])
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))
            for c in col:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        return ws

    sheet("Summary", ["KPI", "Value", "Detail", "Threshold", "Status", "Source"],
          [[k.get("label"), k.get("value"), k.get("detail"), k.get("threshold"), k.get("status"), k.get("source")]
           for k in data.get("kpis") or []], first=True)
    ws = wb["Summary"]
    ws.append([])
    ws.append(["Instance", data.get("instance")])
    ws.append(["Generated", data.get("generated_at")])
    ws.append(["Target endpoints", data.get("target"), data.get("target_source")])
    sheet("Alerts", ["Level", "Message"], [[a.get("level"), a.get("message")] for a in data.get("alerts") or []])
    sheet("Regions", ["Region", "Deployed", "Registered", "Failed", "Fresh %", "Success %", "Health"],
          [[r.get("region"), r.get("deployed"), r.get("registered"), r.get("failed"), r.get("fresh_pct"),
            r.get("success_pct"), r.get("health")] for r in data.get("regions") or []])
    sheet("OS split", ["OS", "Active CIs"], [[o.get("os"), o.get("count")] for o in data.get("os_split") or []])
    sheet("Waves", ["Wave", "Name", "Start", "End", "Regions", "Target", "Deployed", "Status", "Blockers"],
          [[w.get("wave_id"), w.get("name"), w.get("start"), w.get("end"), w.get("regions"), w.get("target"),
            w.get("deployed"), w.get("status"), w.get("blockers")] for w in data.get("waves") or []])
    sheet("Stage gates", ["Gate", "Name", "Unlocks", "Due", "Decision", "Approver", "Decided (UTC)", "Reason"],
          [[g.get("gate_id"), g.get("name"), g.get("unlocks_wave"), g.get("due"), g.get("status"),
            g.get("approver_identity"), g.get("decided_at"), g.get("reason")] for g in data.get("gates") or []])
    sheet("Brand tiers", ["Tier", "Brands", "Target", "Deployed (live)", "Health", "Issues"],
          [[b.get("tier"), b.get("brands"), b.get("target"), b.get("deployed_live"), b.get("health"), b.get("issues")]
           for b in data.get("brand_tiers") or []])
    sheet("Duplicate serials", ["Serial number", "Active CIs"],
          [[d.get("serial_number"), d.get("count")] for d in (data.get("duplicates") or {}).get("serial_groups") or []])
    br = body.get("briefing")
    if isinstance(br, dict):
        rows = [["Confidence", br.get("confidence")], ["Risk trend", br.get("risk_trend")],
                ["Executive summary", br.get("executive_summary")], ["GO/NO-GO (advisory)", br.get("go_no_go")]]
        rows += [["Risk", f"[{r.get('severity')}] {r.get('risk')} — {r.get('evidence', '')}"] for r in br.get("key_risks") or []
                 if isinstance(r, dict)]
        rows += [["Action", f"{a.get('action')} — {a.get('owner', '')} ({a.get('due', '')})"] for a in br.get("actions") or []
                 if isinstance(a, dict)]
        rows += [["Next 48h", x] for x in br.get("next_48h") or []]
        rows += [["Human approval", x] for x in br.get("human_approval_items") or []]
        sheet("AI briefing", ["Item", "Detail"], rows)
    buf = io.BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="ACC_Command_Center_Report.xlsx"'})
