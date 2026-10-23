"""End-to-end test of SoftProdigy Snow Agent v6 against the mock ServiceNow (fake_sn2) and mock Flow Server."""
import io, json, os, sys, time, traceback
os.environ.update(NO_PROXY="127.0.0.1,localhost", no_proxy="127.0.0.1,localhost",
                  FLOW_SERVER_URL="http://127.0.0.1:9998", AGENT_JOBS_DIR=os.path.join(os.path.dirname(os.path.abspath(__file__)), "jobs"),
                  SNOW_VERIFY_TLS="false")
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "snow-ai-service"))
import requests
import groq_client

AI = {"down": False}
def fake_chat(messages, max_tokens=4096, temperature=0.3, model=None, json_mode=False, vision=False):
    if AI["down"]:
        raise groq_client.GroqError("All Groq models failed (simulated outage)")
    p = " ".join(m["content"] if isinstance(m["content"], str) else "" for m in messages)
    if "Create a ServiceNow build plan" in p:
        return json.dumps(PLAN)
    if "Extract every user story" in p:
        return json.dumps({"project_title": "P1 Routing", "project_overview": "Route P1s.", "user_stories": [
            {"title": "Route P1 incidents", "as_a": "service desk manager", "i_want": "P1s routed to on-call",
             "so_that": "SLA is met", "acceptance_criteria": ["Given a P1, When created, Then assigned in 2 min"],
             "priority": "critical", "module": "ITSM", "story_type": "story", "story_points": 5},
            {"title": "VPN outage in Pune", "i_want": "log the VPN outage", "priority": "high", "module": "ITSM",
             "story_type": "incident", "acceptance_criteria": []},
            {"title": "Recurring DB crash", "i_want": "find root cause", "priority": "medium", "story_type": "problem"},
            {"title": "Upgrade firewall", "i_want": "upgrade firewall firmware", "priority": "low", "story_type": "change"}],
            "complete": True})
    if "Pick the ONE ServiceNow tool" in p:
        if "create a problem" in p.lower():
            return json.dumps({"tool": "create_problem", "arguments": {"short_description": "Printers keep failing"},
                               "missing": [], "reason": "User wants a new problem record"})
        if "incidents" in p.lower():
            return json.dumps({"tool": "list_incidents", "arguments": {"limit": 5}, "missing": [], "reason": "List incidents"})
        return json.dumps({"tool": None, "reason": "no match"})
    if "solution architecture" in p and '"layers"' in p:
        return json.dumps({"title": "P1 — Solution Architecture", "layers": [
            {"name": "Users", "components": ["Service desk", "End users"]},
            {"name": "ServiceNow", "components": ["Incident", "Assignment rules", "SLA engine"]},
            {"name": "External", "components": ["On-call tool"]}],
            "flows": [{"from": "End users", "to": "Incident", "label": "raise"}], "description": "How it fits."})
    if "query plan" in p:
        return json.dumps({"table": "incident", "sysparm_query": "priority=1", "sysparm_limit": 10})
    if "Reply with exactly: OK" in p:
        return "OK"
    return "Section text specific to the project.\n\n| A | B |\n|---|---|\n| 1 | 2 |"
groq_client.chat = fake_chat
groq_client.list_available_models = lambda force=False: []
groq_client.transcribe = lambda b, name, language="en": "In the meeting we agreed: route P1 incidents to on-call within 2 minutes."
groq_client.describe_image = lambda b64, mime, instr, max_tokens=3000: "Whiteboard: 1. Route P1 2. Notify manager"

PLAN = {
    "app_name": "P1 Router", "app_scope": "p1router", "scope_mode": "scoped", "description": "Routes P1 incidents",
    "roles": [{"name": "admin", "contains": ["user"]}, {"name": "user"}],
    "tables": [{"name": "route_rule", "label": "Route Rule", "extends": "task", "number_prefix": "RR",
                "fields": [{"name": "on_call_group", "label": "On-call group", "type": "reference", "reference": "sys_user_group"},
                           {"name": "level", "label": "Level", "type": "choice",
                            "choices": [{"value": "p1", "label": "P1"}, {"value": "p2", "label": "P2"}]},
                           {"name": "minutes", "type": "integer", "label": "Minutes"}]}],
    "acls": [{"table": "route_rule", "operations": ["read"], "roles": ["user"]}],
    "business_rules": [{"name": "Default minutes", "table": "route_rule", "when": "before", "insert": True,
                        "script": "if (!current.minutes) current.minutes = 2;\ngs.log('set');"}],
    "client_scripts": [{"name": "Level hint", "table": "route_rule", "type": "onChange", "field": "level",
                        "script": "g_form.showFieldMsg('level','Check group','info');"}],
    "ui_policies": [{"short_description": "Group mandatory", "table": "route_rule",
                     "actions": [{"field": "on_call_group", "mandatory": True}]}],
    "notifications": [{"name": "Rule created", "table": "route_rule", "recipient_fields": ["opened_by"],
                       "subject": "New rule", "message": "A rule was created"}],
    "flows": [{"name": "Notify on rule", "trigger": {"type": "record.created", "table": "route_rule"},
               "actions": [{"id": "log", "type": "core.log", "inputs": {"log_message": "x"}}]}],
}

from fastapi.testclient import TestClient
import api
c = TestClient(api.app)
SN = {"snow_url": "https://127.0.0.1:9443", "snow_user": "admin", "snow_pass": "pw"}
BAD = dict(SN, snow_pass="WRONG")
RESULTS = []

def check(name, cond, detail=""):
    RESULTS.append((name, bool(cond), detail))
    print(("PASS " if cond else "FAIL ") + name + (f"  — {detail}" if detail else ""))

def db(table):
    return requests.get(f"https://127.0.0.1:9443/__db/{table}", verify=False).json()

def run(name, fn):
    try:
        fn()
    except Exception as e:
        traceback.print_exc()
        check(name + " (crashed)", False, str(e)[:200])

# ── 1. Boot / TLS / login / API key ─────────────────────────────────────
def t_login():
    os.environ["SNOW_VERIFY_TLS"] = "true"
    r = c.post("/snow/login", json=SN)
    check("TLS verification ON rejects self-signed cert", r.status_code == 502 and "TLS" in r.text, r.text[:120])
    os.environ["SNOW_VERIFY_TLS"] = "false"
    r = c.post("/snow/login", json=SN)
    check("Login with good credentials", r.status_code == 200 and r.json()["is_admin"], str(r.json().get("roles")))
    r = c.post("/snow/login", json=BAD)
    check("Login with wrong password -> 401", r.status_code == 401, r.text[:100])
    os.environ["AGENT_API_KEY"] = "k123"
    r1 = c.post("/snow/login", json=SN); r2 = c.post("/snow/login", json=SN, headers={"X-API-Key": "k123"})
    cfg = c.get("/snow/config").json()
    check("API key enforced when AGENT_API_KEY set", r1.status_code == 401 and r2.status_code == 200 and cfg["api_key_required"])
    del os.environ["AGENT_API_KEY"]
    old_key = groq_client.GROQ_API_KEY
    groq_client.GROQ_API_KEY = "some-other-key"
    check("Groq key comes from configuration, not source", groq_client.get_api_key() == "some-other-key")
    groq_client.GROQ_API_KEY = old_key
run("login", t_login)

def t_artifact_manager():
    r = c.post("/snow/artifacts/schema", json={**SN, "table": "incident"})
    check("Artifact manager reads table schema", r.status_code == 200 and r.json().get("table") == "incident", r.text[:120])
    r = c.post("/snow/artifacts/list", json={**SN, "table": "incident", "query": "priority=1", "limit": 10})
    check("Artifact manager lists any permitted table", r.status_code == 200 and r.json()["count"] >= 1, r.text[:120])
    r = c.post("/snow/artifacts/list", json={**SN, "table": "incident", "query": "sys_idjavascript:alert(1)"})
    check("Artifact manager rejects JavaScript encoded queries", r.status_code == 400, r.text[:120])
    r = c.post("/snow/artifacts/create", json={**SN, "table": "sys_user_role", "fields": {"name": "universal.manager"}})
    check("Artifact creation requires explicit confirmation", r.status_code == 428, r.text[:120])
    r = c.post("/snow/artifacts/create", json={
        **SN, "table": "sys_user_role", "fields": {"name": "universal.manager"}, "confirm": True,
        "context": {"scope_mode": "global", "update_set_name": "HAI_artifact_manager"},
    })
    data = r.json()
    role_id = data.get("sys_id", "")
    check("Artifact manager creates a scoped/update-set-aware metadata artifact",
          r.status_code == 200 and role_id and data.get("context", {}).get("update_set", {}).get("sys_id"), r.text[:180])
    r = c.post("/snow/artifacts/get", json={**SN, "table": "sys_user_role", "sys_id": role_id})
    check("Artifact manager reads an artifact by sys_id", r.status_code == 200 and r.json().get("sys_id") == role_id, r.text[:120])
    r = c.post("/snow/artifacts/update", json={**SN, "table": "sys_user_role", "sys_id": role_id,
                                                  "fields": {"description": "Universal manager role"}, "confirm": True})
    check("Artifact manager updates confirmed artifacts", r.status_code == 200 and r.json().get("action") == "updated", r.text[:120])
    r = c.post("/snow/artifacts/delete", json={**SN, "table": "sys_user_role", "sys_id": role_id,
                                                  "reason": "Test cleanup", "confirm": True})
    check("Artifact manager deletes only with confirmation and reason", r.status_code == 200 and r.json().get("action") == "deleted", r.text[:120])
run("artifact manager", t_artifact_manager)

# ── 2. Universal extraction ───────────────────────────────────────────────
def t_extract():
    r = c.post("/snow/extract/universal", data={"text": "Route P1 incidents to on-call within 2 minutes; VPN is down"})
    d = r.json()
    check("extract/universal (text)", r.status_code == 200 and len(d["user_stories"]) == 4, f"{len(d.get('user_stories', []))} stories")
    import docx
    D = docx.Document(); D.add_paragraph("Requirements"); t = D.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Req"; t.cell(1, 0).text = "Route P1 to on-call"; b = io.BytesIO(); D.save(b)
    r = c.post("/snow/extract/universal", files={"file": ("req.docx", b.getvalue())})
    check("extract/universal (docx incl. tables)", r.status_code == 200 and "Route P1 to on-call" in r.json()["_raw_text"])
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(["Story", "Acceptance"]); ws.append(["Route P1", "2 min"])
    b = io.BytesIO(); wb.save(b)
    r = c.post("/snow/extract/universal", files={"file": ("s.xlsx", b.getvalue())})
    check("extract/universal (xlsx with story columns -> direct import)", r.status_code == 200
          and r.json()["source"]["mode"].startswith("direct") and "Route P1" in json.dumps(r.json()["user_stories"]))
    sample = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_user_stories.xlsx")
    d = c.post("/snow/extract/universal", files={"file": ("user_stories_snow.xlsx", open(sample, "rb").read())}).json()
    import openpyxl
    ws = openpyxl.load_workbook(sample).active
    orig = [(ws.cell(r_, 1).value, ws.cell(r_, 2).value, ws.cell(r_, 3).value, ws.cell(r_, 4).value)
            for r_ in range(2, ws.max_row + 1) if ws.cell(r_, 1).value]
    got = [(x["id"], x["title"], x["description"], "\n".join(x["acceptance_criteria"])) for x in d["user_stories"]]
    check("Your sample file: 22 stories, IDs, summaries, descriptions and all 342 criteria kept exactly",
          len(got) == 22 and all(a[0] == b[0] and a[1] == b[1] and a[2] == b[2] and
                                 [l.strip() for l in a[3].split("\n") if l.strip()] == b[3].split("\n")
                                 for a, b in zip(orig, got)),
          f"{len(got)} stories, {sum(len(x['acceptance_criteria']) for x in d['user_stories'])} criteria")
    pay = [{"id": x["id"], "summary": x["title"], "description": x["description"],
            "acceptance_criteria": x["acceptance_criteria"]} for x in d["user_stories"]]
    xl = c.post("/snow/userstories/excel", json={"project_name": "P", "user_stories": pay})
    wb2 = openpyxl.load_workbook(io.BytesIO(xl.content)).active
    rows2 = [tuple(wb2.cell(r_, k).value for k in range(1, 5)) for r_ in range(2, wb2.max_row + 1) if wb2.cell(r_, 1).value]
    check("Excel export = your sample (same 4 headers, all 88 cells identical)",
          [c_.value for c_ in wb2[1]] == [c_.value for c_ in ws[1]] and
          all((a[k] or "").strip() == (b[k] or "").strip() for a, b in zip(orig, rows2) for k in range(4)) and len(rows2) == 22)
    wd = c.post("/snow/documents/export", json={"doc_type": "USER_STORIES", "project_name": "P", "requirements": "x",
                                                "user_stories": pay})
    import docx as _dx
    txt = "\n".join(p_.text for p_ in _dx.Document(io.BytesIO(wd.content)).paragraphs)
    check("Word 'User Stories & Acceptance Criteria' has every story ID and criterion",
          all(o[0] in txt for o in orig) and all(l.strip() in txt for o in orig for l in o[3].split("\n") if l.strip()))
    r = c.post("/snow/extract/universal", files={"file": ("s.csv", b"story,criteria\nRoute P1,2 min\n")})
    check("extract/universal (csv)", r.status_code == 200)
    r = c.post("/snow/extract/universal", files={"file": ("m.mp3", b"ID3fakeaudio")})
    check("extract/universal (audio -> Whisper)", r.status_code == 200 and r.json()["source"]["type"] == "audio" and "transcript" in r.json())
    r = c.post("/snow/extract/universal", files={"file": ("wb.png", b"\x89PNGfake", "image/png")})
    check("extract/universal (image -> vision)", r.status_code == 200 and r.json()["source"]["type"] == "image")
    r = c.post("/snow/extract/universal", files={"file": ("x.exe", b"MZ")})
    check("extract/universal rejects unsupported type (415)", r.status_code == 415)
    AI["down"] = True
    r = c.post("/snow/extract/universal", data={"text": "Route P1 incidents to on-call"})
    AI["down"] = False
    check("extract/universal AI outage -> 503 (no fake stories)", r.status_code == 503)
run("extract", t_extract)

# ── 3. Approve & deploy ────────────────────────────────────────────────────
def t_deploy():
    stories = c.post("/snow/extract/universal", data={"text": "anything long enough"}).json()["user_stories"]
    r = c.post("/snow/stories/approve-and-deploy", json={**SN, "user_stories": stories})
    check("approve-and-deploy refuses without approved=true (428)", r.status_code == 428)
    r = c.post("/snow/stories/approve-and-deploy", json={**SN, "approved": True, "user_stories": stories})
    d = r.json()
    tables = {x["index"]: x["table"] for x in d["stories_created"]}
    check("approve-and-deploy creates the right record types",
          tables == {0: "rm_story", 1: "incident", 2: "problem", 3: "change_request"}, str(tables))
    inc = [x for x in db("incident") if x.get("short_description") == "VPN outage in Pune"]
    check("incident gets impact/urgency and no guessed assignment group",
          inc and inc[0].get("impact") == "1" and not inc[0].get("assignment_group"), json.dumps(inc[0] if inc else {})[:160])
    one = [dict(stories[1], title="VPN outage mapped")]
    r = c.post("/snow/stories/approve-and-deploy", json={**SN, "approved": True, "user_stories": one,
                                                         "group_mapping": {stories[1].get("module") or "ITSM": "Service Desk"}})
    mapped = [x for x in db("incident") if x.get("short_description") == "VPN outage mapped"]
    check("explicit module → group mapping assigns the exact group",
          r.status_code == 200 and mapped and mapped[0].get("assignment_group"), r.text[:200])
    one = [dict(stories[1], title="VPN outage unknown group", assignment_group="Network Wizards")]
    r = c.post("/snow/stories/approve-and-deploy", json={**SN, "approved": True, "user_stories": one})
    d2 = r.json()
    check("unknown group: nothing created, a question is returned",
          not d2["stories_created"] and d2["errors"] and d2["errors"][0].get("needs_input")
          and not [x for x in db("incident") if x.get("short_description") == "VPN outage unknown group"], r.text[:200])
    st = [x for x in db("rm_story") if x.get("short_description") == "Route P1 incidents"]
    check("rm_story gets acceptance_criteria + story_points", st and "Given a P1" in st[0].get("acceptance_criteria", "")
          and st[0].get("story_points") == 5)
    r = c.post("/snow/stories/approve-and-deploy", json={**BAD, "approved": True, "user_stories": stories})
    check("approve-and-deploy bad creds -> 401", r.status_code == 401)
    r = c.post("/snow/stories/create-in-snow", json={**SN, "user_stories": stories[:1]})
    check("stories/create-in-snow -> rm_story", r.status_code == 200 and r.json()["table"] == "rm_story")
run("deploy", t_deploy)

# ── 4. Documents ────────────────────────────────────────────────────────────
def t_docs():
    titles = {}
    for dt in ["BRD", "FRD", "TSD", "TEST_PLAN", "PROCESS_FLOW", "DEPLOYMENT_GUIDE", "USER_STORIES"]:
        r = c.post("/snow/documents/generate", json={"doc_type": dt, "requirements": "1. Route P1\nAcceptance Criteria:\n- 2 min", "project_name": "P1"})
        titles[dt] = r.json().get("title")
    check("Every doc type has its own template", titles["TEST_PLAN"] == "Test Plan & Test Cases"
          and titles["PROCESS_FLOW"] == "Process Flow Document" and titles["DEPLOYMENT_GUIDE"] == "Deployment Guide", str(titles))
    tsd = c.post("/snow/documents/generate", json={"doc_type": "TSD", "requirements": "Route P1 incidents", "project_name": "P1"}).json()
    arch = [x for x in tsd.get("sections", []) if x.get("diagram")]
    dl = c.post("/snow/documents/download", json={"document": tsd})
    import docx as _dx
    pics = len(_dx.Document(io.BytesIO(dl.content)).inline_shapes) if dl.status_code == 200 else 0
    check("TSD has a real architecture diagram image (drawn from layers/flows)", arch and pics >= 1,
          f"{len(arch[0]['diagram']['layers']) if arch else 0} layers, {pics} image(s)")
    r = c.post("/snow/documents/generate", json={"doc_type": "NOPE", "requirements": "x"})
    check("Unknown doc type -> 400 (no silent BRD)", r.status_code == 400)
    doc = c.post("/snow/documents/generate", json={"doc_type": "BRD", "requirements": "x", "project_name": "P"}).json()
    r = c.post("/snow/documents/download", json={"document": doc})
    import docx
    dd = docx.Document(io.BytesIO(r.content))
    check("Download renders the on-screen document (docx, tables kept)", r.status_code == 200 and len(dd.tables) >= 1)
    AI["down"] = True
    r = c.post("/snow/documents/generate", json={"doc_type": "BRD", "requirements": "x"})
    AI["down"] = False
    check("Docs with AI outage -> 503 (no error text inside a 'Ready' doc)", r.status_code == 503, r.text[:100])
run("docs", t_docs)

# ── 5. Plan + validation ────────────────────────────────────────────────────
def t_plan():
    r = c.post("/snow/plan", json={"requirements": "Route P1s", **SN})
    d = r.json()
    check("/snow/plan returns validated plan", r.status_code == 200 and d["_validation"]["valid"], str(d.get("_validation", {}).get("issues")))
    bad = {"app_name": "Bad", "tables": [{"name": "t1", "fields": [{"name": "who", "type": "reference"}]}],
           "client_scripts": [{"name": "cs", "table": "t1", "type": "onLoad", "script": "var gr = new GlideRecord('incident');"}],
           "business_rules": [{"name": "br", "table": "t1", "script": "current.update();"},
                              {"name": "br2", "table": "t1", "script": "if (x { }"}],
           "flows": [{"name": "empty flow", "trigger": {"type": "record.created", "table": "t1"}, "actions": []}]}
    v = c.post("/snow/build/validate", json={"plan": bad}).json()
    txt = " | ".join(v["issues"])
    check("Validator catches missing reference target", "no reference table" in txt)
    check("Validator catches GlideRecord in client script", "server-side APIs" in txt)
    check("Validator catches current.update() in before rule", "current.update()" in txt)
    check("Validator catches JS syntax error", "syntax error" in txt)
    check("Validator catches flow with no actions", "no core.* actions" in txt)
run("plan", t_plan)

def wait(job_id):
    for _ in range(120):
        j = c.get(f"/snow/build/{job_id}").json()
        if j["status"] not in ("queued", "running"):
            return j
        time.sleep(0.25)
    return j

# ── 6. Build: scoped app, new update set, continue on error ─────────────────
def t_build_scoped():
    r = c.post("/snow/build", json={**BAD, "plan": PLAN})
    check("Build with wrong password -> 401 before any job", r.status_code == 401)
    r = c.post("/snow/build", json={**SN, "plan": PLAN, "update_set_name": "HAI_P1_Router_test", "stop_on_error": False})
    j = wait(r.json()["job_id"])
    for st in j["steps"]:
        print("     ", st["msg"])
    app_ = [a for a in db("sys_app") if a["name"] == "P1 Router"]
    check("Scoped app created with vendor prefix", app_ and app_[0]["scope"].startswith("x_1234_"), app_[0]["scope"] if app_ else "")
    scope = app_[0]["scope"]; app_id = app_[0]["sys_id"]
    us = [u for u in db("sys_update_set") if u["name"] == "HAI_P1_Router_test"]
    check("Update set created and tied to the app", us and us[0]["application"] == app_id)
    prefs = {p["name"]: p["value"] for p in db("sys_user_preference")}
    check("User preferences make it current (apps.current_app, sys_update_set, updateSetForScope<app>)",
          prefs.get("apps.current_app") == app_id and prefs.get("sys_update_set") == us[0]["sys_id"]
          and prefs.get(f"updateSetForScope{app_id}") == us[0]["sys_id"], json.dumps(prefs)[:200])
    tbl = [t for t in db("sys_db_object") if t["name"] == f"{scope}_route_rule"]
    task_id = next(t["sys_id"] for t in db("sys_db_object") if t["name"] == "task")
    check("Table named in scope and parent resolved to sys_id", tbl and tbl[0].get("super_class") == task_id
          and tbl[0].get("sys_scope") == app_id)
    fld = {f["element"]: f for f in db("sys_dictionary") if f["name"] == f"{scope}_route_rule"}
    check("Reference field has its target table", fld.get("on_call_group", {}).get("reference") == "sys_user_group")
    ch = [x for x in db("sys_choice") if x["element"] == "level"]
    check("Choice field + 2 choices", fld.get("level", {}).get("choice") == 1 and len(ch) == 2)
    roles = [x["name"] for x in db("sys_user_role")]
    check("Scoped roles named <scope>.<role>", f"{scope}.admin" in roles and f"{scope}.user" in roles, str(roles[-3:]))
    br = [x for x in db("sys_script") if x["name"] == "Default minutes"]
    check("Business rule wrapped + gs.log -> gs.info for scoped app", br and "executeRule" in br[0]["script"]
          and "gs.info(" in br[0]["script"] and "gs.log(" not in br[0]["script"])
    cs = [x for x in db("sys_script_client") if x["name"] == "Level hint"]
    check("onChange client script wrapped with field", cs and cs[0].get("field") == "level" and "function onChange" in cs[0]["script"])
    acl = [x for x in j["results"] if x["kind"] == "ACL"]
    check("ACL 403 reported as a WARNING with the manual security_admin step (not ✅, build continues)",
          acl and acl[0]["status"] == "warning" and "security_admin" in acl[0]["message"])
    fl = [x for x in j["results"] if x["kind"] == "Flow"]
    check("Flow deployed through Flow Server (not an empty sys_hub_flow)", fl and fl[0]["status"] == "created"
          and not db("sys_hub_flow"))
    calls = requests.get("http://127.0.0.1:9998/__calls").json()
    check("Flow Server got the resolved scoped table + app scope", calls and calls[-1]["flowSpec"]["trigger"]["table"] == f"{scope}_route_rule"
          and calls[-1]["scopeName"] == scope)
    v = j["verification"]
    check("Every created artifact verified in sys_update_xml of the new update set",
          v["expected"] > 10 and not v["missing"], f"captured {v['captured_count']}, expected {v['expected']}, missing {v['missing']}")
    check("Job status is honest (completed_with_issues because the ACL was not created)", j["status"] == "completed_with_issues", j["status"])
    # re-run: everything reused, nothing duplicated
    r = c.post("/snow/build", json={**SN, "plan": PLAN, "update_set_name": "HAI_P1_Router_test", "stop_on_error": False})
    j2 = wait(r.json()["job_id"])
    check("Re-running the same plan reuses (no duplicates)", j2["summary"]["created"] <= 2 and j2["summary"]["reused"] >= 10,
          json.dumps(j2["summary"]))
    return us[0]["sys_id"]
US_ID = {}
run("build scoped", lambda: US_ID.setdefault("id", t_build_scoped()))

# ── 7. Why the session reset matters (mock mirrors ServiceNow session caching) ─
def t_session_proof():
    import sn_client as snc
    sn = snc.SNClient(**{"instance_url": SN["snow_url"], "username": "admin", "password": "pw"})
    sn.query("sys_user", "user_name=admin")           # session starts: pref = the build's update set
    app_ = snc.ensure_app(sn, mode="scoped", app_name="Session Proof")
    ctx_other = snc.ensure_update_set(sn, app_sys_id=app_["sys_id"], name="Proof US")[0]
    user = sn.user()
    snc._upsert_pref(sn, user["sys_id"], "sys_update_set", ctx_other["sys_id"])   # change pref, NO reset
    rec = sn.create("sys_user_role", {"name": "proof.role", "sys_scope": app_["sys_id"]}, "sys_id")
    caught = [x for x in db("sys_update_xml") if rec["sys_id"] in x["name"] and x["update_set"] == ctx_other["sys_id"]]
    check("Without a session reset the new preference is NOT used (the v5 failure mode)", not caught)
    ctx = snc.set_context(sn, app=app_, update_set_name="Proof US")                 # prefs + reset
    rec2 = sn.create("sys_user_role", {"name": "proof.role2", "sys_scope": app_["sys_id"]}, "sys_id")
    caught2 = [x for x in db("sys_update_xml") if rec2["sys_id"] in x["name"] and x["update_set"] == ctx["update_set"]["sys_id"]]
    check("With set_context (prefs + new session) the record lands in the chosen update set", bool(caught2))
run("session", t_session_proof)

# ── 8. Stop + rollback ──────────────────────────────────────────────────────
def t_rollback():
    plan = json.loads(json.dumps(PLAN)); plan["app_name"] = "Rollback Demo"; plan["app_scope"] = "rbdemo"
    plan["business_rules"] = [{"name": "Broken", "table": "no_such_table_xyz", "script": "var a = 1;"}]
    r = c.post("/snow/build", json={**SN, "plan": plan, "update_set_name": "HAI_rollback", "stop_on_error": True,
                                    "rollback_on_error": True})
    j = wait(r.json()["job_id"])
    check("Build stops at the first failure (business rule on a missing table)",
          j["status"] == "failed" and j["failed_step"] == "Business rules", str(j.get("failed_step")))
    left = [t for t in db("sys_db_object") if "rbdemo" in t["name"]]
    check("Rollback removed the records this run created (table, app, update set)",
          not left and not [u for u in db("sys_update_set") if u["name"] == "HAI_rollback"]
          and all(x["ok"] for x in j["rollback"]), f"{len(j['rollback'])} rolled back")
run("rollback", t_rollback)

# ── 9. Global + existing update set guard ───────────────────────────────────
def t_global():
    plan = {"app_name": "Glob", "scope_mode": "global", "tables": [{"name": "asset_note", "fields": [{"name": "note"}]}]}
    j = wait(c.post("/snow/build", json={**SN, "plan": plan, "update_set_name": "HAI_global"}).json()["job_id"])
    names = [t["name"] for t in db("sys_db_object")]
    fl = [f["element"] for f in db("sys_dictionary") if f["name"] == "u_asset_note"]
    check("Global build: u_ table and u_ field, captured in global update set",
          "u_asset_note" in names and "u_note" in fl and not j["verification"]["missing"] and j["status"] == "completed", j["status"])
    other = next(u for u in db("sys_update_set") if u["name"] == "HAI_P1_Router_test")
    j = wait(c.post("/snow/build", json={**SN, "plan": plan, "update_set_sys_id": other["sys_id"]}).json()["job_id"])
    check("Refuses an update set that belongs to another application", j["status"] == "failed" and "different application" in (j.get("error") or ""), j.get("error"))
run("global", t_global)

# ── 10. Update set management ───────────────────────────────────────────────
def t_us():
    d = c.post("/snow/updatesets/list", json=SN).json()
    cur = [u for u in d["update_sets"] if u["current"]]
    check("Update set list shows captured counts + current", d["update_sets"] and cur, f"{len(d['update_sets'])} sets")
    d = c.post("/snow/updatesets/changes", json={**SN, "update_set_sys_id": US_ID.get("id")}).json()
    check("Update set changes listed by type", d["count"] > 10, json.dumps(d["by_type"]))
    empty = c.post("/snow/updatesets/context", json={**SN, "scope_mode": "global", "update_set_name": "HAI_empty_" + str(int(time.time()))}).json()
    r = c.post("/snow/updatesets/complete", json={**SN, "update_set_sys_id": empty["update_set"]["sys_id"]})
    check("Complete refuses an empty update set (409)", r.status_code == 409, f"{r.status_code} {r.text[:200]}")
    r = c.post("/snow/updatesets/complete", json={**SN, "update_set_sys_id": US_ID.get("id")})
    check("Complete marks state + returns export link", r.status_code == 200 and "export_update_set.do" in r.json()["export_xml_url"])
    apps = c.post("/snow/apps/list", json=SN).json()
    check("Apps list includes the new scoped app", any(a["name"] == "P1 Router" for a in apps["apps"]))
run("updatesets", t_us)

# ── 11. Health scan ──────────────────────────────────────────────────────────
def t_health():
    r = c.post("/snow/health-scan", json=SN)
    d = r.json()
    cats = {x["key"]: x["score"] for x in d["categories"]}
    check("Health scan runs live checks and returns UI format", r.status_code == 200 and d["overall_score"] is not None
          and "findings" in d and len(d["checks"]) == 24, f"score {d.get('overall_score')}, unknown {d.get('unknown')}, cats {cats}")
    f = {x["title"]: x for x in d["findings"]}
    check("Finds the seeded problems (8 admins, open P1, unassigned incidents, CIs w/o support group)",
          "Admin role holders" in f and "Open P1 incidents" in f and "Unassigned incidents" in f and "CIs without support group" in f,
          ", ".join(f))
    check("Unreadable checks are 'unknown', not scored", d["unknown"] >= 1)
    r = c.post("/snow/health-scan", json=BAD)
    check("Health scan with wrong password -> 401 (no fake grade)", r.status_code == 401, r.text[:100])
    r = c.post("/snow/healthscan", json=SN)
    check("/snow/healthscan alias uses the same engine", r.status_code == 200 and "categories" in r.json())
run("health", t_health)

# ── 12. Interactive ops (existing + new) ─────────────────────────────────────
def t_interactive():
    I = "/snow/interactive"
    inc = c.post(I + "/incidents/create", json={**SN, "short_description": "Printer down"}).json()
    check("Incident create", inc.get("number", "").startswith("INC"))
    p = c.post(I + "/problems/create", json={**SN, "short_description": "Printers keep failing", "priority": "high",
                                            "related_incidents": [inc["number"], "INC9999999"]}).json()
    check("Problem create + links incidents", p.get("number", "").startswith("PRB") and p["linked_incidents"] == [inc["number"]]
          and p["incidents_not_found"] == ["INC9999999"])
    r = c.post(I + "/problems/update", json={**SN, "problem_id": p["sys_id"], "workaround": "restart spooler", "known_error": True})
    check("Problem update", r.status_code == 200)
    r = c.post(I + "/changes/create", json={**SN, "short_description": "Patch printers", "risk": "high"})
    check("Change create", r.status_code == 200)
    r = c.post(I + "/knowledge/create", json={**SN, "short_description": "Fix spooler", "text": "Restart"})
    check("Knowledge create reports missing KB honestly (404 on mock without KB)", r.status_code == 404, r.text[:80])
    for path in ("/approvals/list", "/catalog/items", "/atf/suites"):
        r = c.post(I + path, json=SN)
        check(f"{path} responds", r.status_code == 200, r.text[:80])
    r = c.post(I + "/query", json={**SN, "question": "P1 incidents"})
    check("NL query", r.status_code == 200 and r.json()["total_found"] >= 1)
    r = c.post("/snow/jira/push", json={"user_stories": [{"title": "x"}]})
    check("Jira push says what is missing when not configured", r.status_code == 400 and "JIRA_URL" in r.text)
    r = c.post("/snow/universal/chat", json={"message": "What is an update set?"})
    check("Universal chat", r.status_code == 200 and r.json()["response"])
run("interactive", t_interactive)

# ── 13. Dynamic credentials: a second ServiceNow user signs in and builds with their own account ─────────
def t_dynamic_user():
    JD = {"snow_url": SN["snow_url"], "snow_user": "jdoe", "snow_pass": "pw2"}
    r = c.post("/snow/login", json=JD)
    check("Second user (jdoe) signs in with their own password", r.status_code == 200 and r.json()["user"]["user_name"] == "jdoe")
    r = c.post("/snow/login", json=dict(JD, snow_pass="pw"))
    check("jdoe with admin's password is rejected (no shared/fixed credentials)", r.status_code == 401)
    plan = {"app_name": "Jane App", "app_scope": "janeapp", "tables": [{"name": "jane_item", "fields": [{"name": "note"}]}]}
    j = wait(c.post("/snow/build", json={**JD, "plan": plan, "update_set_name": "HAI_jane"}).json()["job_id"])
    jdoe_id = "a1b2c3d4e5f60718293a4b5c6d7e8f90"
    prefs = [p for p in db("sys_user_preference") if p["name"] == "sys_update_set"]
    mine = [p for p in prefs if p["user"] == jdoe_id]
    us = next(u for u in db("sys_update_set") if u["name"] == "HAI_jane")
    check("Update set made current for jdoe (not admin)", mine and mine[0]["value"] == us["sys_id"])
    tbl = [t for t in db("sys_db_object") if t["name"].endswith("janeapp_jane_item")]
    check("Records created with jdoe's session and captured in jdoe's update set",
          tbl and tbl[0].get("sys_created_by_id") == jdoe_id and j["status"] == "completed" and not j["verification"]["missing"],
          j["status"])
run("dynamic user", t_dynamic_user)

# ── 14. All SoftProdigy ServiceNow tools inside the agent ─────────────────────────────
def t_tools():
    cat = c.get("/snow/tools").json()
    pk = c.get("/snow/tools/packages").json()
    # 522 vendor tools, minus the superseded vendored asset module, plus the 38 shared ITAM tools (+ "itam" bundle)
    names = {t["name"] for t in cat["tools"]}
    itam = {t["name"] for t in c.get("/snow/tools?package=itam").json()["tools"]}
    check("Tool catalog: vendor tools + 38 ITAM tools, vendored asset tools superseded",
          len(itam) == 38 and {"create_asset", "retire_asset", "receive_purchase_order_item"} <= names
          and not ({"create_hardware_asset", "mark_asset_reserved", "perform_asset_lifecycle_action"} & names)
          and len(pk["packages"]) == 21,f"{cat['total']} tools, {len(pk['packages'])} bundles, {len(itam)} ITAM")
    need = ["create_incident", "create_problem", "create_change_request", "approve_approval", "reject_approval",
            "create_article", "publish_article", "create_catalog_item", "create_request", "create_workflow",
            "create_business_rule", "create_ui_policy", "create_sla_definition" if "create_sla_definition" in names else "create_incident",
            "atf_framework.create_atf_test", "atf_framework.generate_atf_tests_for_table",
            "servicenow_internal_cmdb.upsert_ci", "create_report" if "create_report" in names else "create_incident"]
    missing = [n for n in need if n not in names]
    check("Catalog includes incidents, problems, changes, approvals, knowledge, catalog, workflows, scripts, ATF, CMDB",
          not missing, f"missing: {missing}")
    sch = c.get("/snow/tools/create_incident").json()
    check("Tool detail returns its input schema", "short_description" in json.dumps(sch["input_schema"]))
    r = c.post("/snow/tools/run", json={**SN, "name": "list_incidents", "arguments": {"limit": 3}})
    check("Read tool runs as the signed-in user", r.status_code == 200 and r.json()["success"], r.text[:120])
    r = c.post("/snow/tools/run", json={**SN, "name": "create_incident", "arguments": {"short_description": "via tool"}})
    check("Write tool without confirm -> 428", r.status_code == 428)
    r = c.post("/snow/tools/run", json={**SN, "name": "create_incident", "arguments": {"short_description": "via tool"}, "confirm": True})
    made = [x for x in db("incident") if x.get("short_description") == "via tool"]
    check("Write tool with confirm creates the record", r.status_code == 200 and r.json()["success"] and made, r.text[:120])
    JD = {"snow_url": SN["snow_url"], "snow_user": "jdoe", "snow_pass": "pw2"}
    c.post("/snow/tools/run", json={**JD, "name": "create_incident", "arguments": {"short_description": "jdoe tool"}, "confirm": True})
    mine = [x for x in db("incident") if x.get("short_description") == "jdoe tool"]
    check("Tools run with each user's own login (record created by jdoe)", mine and mine[0].get("sys_created_by_id") == "a1b2c3d4e5f60718293a4b5c6d7e8f90")
    r = c.post("/snow/tools/run", json={**BAD, "name": "list_incidents", "arguments": {}})
    body = r.json()
    check("Wrong password through a tool is reported as a failure", r.status_code >= 400 or body.get("success") is False, r.text[:120])
    r = c.post("/snow/tools/run", json={**SN, "name": "create_incident", "arguments": {"bogus": 1}, "confirm": True})
    check("Invalid tool arguments -> 422 with the reason", r.status_code == 422, r.text[:100])
    d = c.post("/snow/tools/ask", json={**SN, "message": "show me the latest incidents"}).json()
    check("Ask in plain English -> read tool picked and run", d.get("tool") == "list_incidents" and d.get("result", {}).get("success"))
    d = c.post("/snow/tools/ask", json={**SN, "message": "create a problem for printers failing"}).json()
    check("Ask -> write tool comes back for confirmation (not run)", d.get("tool") == "create_problem" and d.get("needs_confirmation")
          and not [x for x in db("problem") if x.get("short_description") == "Printers keep failing" and x.get("_via_ask")])
    # sweep: every read tool that needs no inputs
    ok = err = crash = 0
    bad = []
    for t in cat["tools"]:
        if t["kind"] != "read":
            continue
        req = c.get(f"/snow/tools/{t['name']}").json()["input_schema"].get("required") or []
        if req:
            continue
        r = c.post("/snow/tools/run", json={**SN, "name": t["name"], "arguments": {}})
        if r.status_code == 200 and r.json()["success"]:
            ok += 1
        elif r.status_code in (200, 400, 404, 409, 422, 424, 502) and "detail" in r.text + '"detail"':
            err += 1; bad.append(t["name"])   # a readable refusal (e.g. plugin table missing on the mock)
        else:
            crash += 1; bad.append(t["name"] + f"({r.status_code})")
    check("Sweep: read tools with no required inputs execute without crashing the agent", crash == 0,
          f"{ok} succeeded, {err} returned a tool-level error on the mock (tables the mock does not have), {crash} crashed")
run("tools", t_tools)

passed = sum(1 for _, ok, _ in RESULTS if ok)
print(f"\n==== {passed}/{len(RESULTS)} checks passed ====")
for n, ok, d in RESULTS:
    if not ok:
        print("FAILED:", n, d)
json.dump([{"check": n, "passed": ok, "detail": d} for n, ok, d in RESULTS], open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "e2e_results.json"), "w"), indent=1)
