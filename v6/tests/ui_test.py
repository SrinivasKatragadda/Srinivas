"""Drive the v6 UI in Chromium against the agent (mock AI) + mock ServiceNow. Screenshots in ./shots."""
import os, time
from playwright.sync_api import sync_playwright

os.makedirs("shots", exist_ok=True)
errors, results = [], []

def ok(name, cond, detail=""):
    results.append((name, bool(cond)))
    print(("PASS " if cond else "FAIL ") + name + (f" — {detail}" if detail else ""))

with sync_playwright() as p:
    b = p.chromium.launch(args=["--no-proxy-server"])
    pg = b.new_page(viewport={"width": 1440, "height": 950})
    pg.on("pageerror", lambda e: errors.append(str(e)))
    pg.on("console", lambda m: m.type == "error" and errors.append(m.text))
    pg.on("dialog", lambda d: d.accept())
    pg.goto(f"http://127.0.0.1:{os.getenv('AGENT_TEST_PORT', '8011')}/ui/")

    # wrong password first
    pg.wait_for_selector("#f-user")
    time.sleep(1)
    labels = [l.strip().lower() for l in pg.locator(".connect-card .field-label").all_inner_texts()]
    ok("Login screen asks only for username and password (instance comes from the agent, no API key box)",
       labels == ["username", "password"] and pg.locator("#f-url").count() == 0 and pg.locator("#f-key").count() == 0
       and "127.0.0.1:9443" in pg.inner_text(".connect-card"), str(labels))
    pg.screenshot(path="shots/00_login.png")
    pg.fill("#f-user", "admin"); pg.fill("#f-pass", "nope")
    pg.click("#btn-connect"); pg.wait_for_selector(".error-bar", timeout=15000)
    ok("Connect with wrong password shows error, stays on login", "rejected" in pg.inner_text(".error-bar"), pg.inner_text(".error-bar"))
    pg.screenshot(path="shots/01_login_error.png")
    pg.fill("#f-pass", "pw"); pg.click("#btn-connect")
    pg.wait_for_selector("text=Signed in to", timeout=15000)
    ok("Connect validates and signs in", True)
    pg.screenshot(path="shots/02_chat.png")

    # chat: create incident with confirmation
    pg.fill("#chat-input", "create incident for VPN outage in Pune"); pg.keyboard.press("Enter")
    pg.wait_for_selector("#_cyes", timeout=10000); pg.click("#_cyes")
    pg.wait_for_selector("text=Incident created", timeout=15000)
    ok("Chat creates incident after confirmation", True)
    pg.screenshot(path="shots/03_chat_incident.png")

    # implement: upload the real sample file -> 22 stories kept as written
    pg.click("[data-nav=implement]")
    pg.click("#tab-file")
    pg.set_input_files("#file-input", os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_user_stories.xlsx"))
    pg.wait_for_selector("text=Review User Stories", timeout=20000)
    n = pg.locator(".story-row").count()
    ok("Upload your sample .xlsx -> 22 stories with original IDs", n == 22 and "VideoPreparationVersion_12" in pg.inner_text("#mc"), f"{n} stories")
    pg.screenshot(path="shots/04a_sample_import.png")
    with pg.expect_download() as dlw:
        pg.click("#btn-dl-stories-word")
    wpath = dlw.value.path()
    import docx, openpyxl
    wtxt = "\n".join(p_.text for p_ in docx.Document(wpath).paragraphs)
    ws_ = openpyxl.load_workbook(os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_user_stories.xlsx")).active
    ids = [ws_.cell(r, 1).value for r in range(2, ws_.max_row + 1) if ws_.cell(r, 1).value]
    acs = [l.strip() for r in range(2, ws_.max_row + 1) for l in str(ws_.cell(r, 4).value or "").split("\n") if l.strip()]
    ok("Review step: download User Stories Word straight away (all 22 IDs + all 342 criteria inside)",
       all(i in wtxt for i in ids) and all(a in wtxt for a in acs), f"{dlw.value.suggested_filename}, {sum(a in wtxt for a in acs)}/{len(acs)} criteria")
    with pg.expect_download() as dlx:
        pg.click("#btn-dl-stories-xlsx")
    dlx.value.save_as("shots/_stories.xlsx")
    wb_ = openpyxl.load_workbook("shots/_stories.xlsx").active
    ok("Review step: download User Stories Excel straight away (22 rows, your 4 columns)",
       wb_.max_row - 1 == 22 and [c.value for c in wb_[1]] == ["User Story ID", "Summary", "Description", "Acceptance Criteria"])
    with pg.expect_download() as dl0:
        pg.click("#btn-gen-docs"); pg.wait_for_selector("[data-dl='0'][data-fmt='xlsx']", timeout=60000)
        pg.click("[data-dl='0'][data-fmt='xlsx']")
    ok("User Stories Excel downloads from the screen", dl0.value.suggested_filename.endswith(".xlsx"), dl0.value.suggested_filename)
    pg.wait_for_selector("text=/documents generated/", timeout=60000)
    pg.click("#btn-restart") if pg.locator("#btn-restart").count() else None
    pg.click("[data-nav=implement]")
    if pg.locator("#btn-back-1").count(): pg.click("#btn-back-1")
    if pg.locator("#btn-back-0").count(): pg.click("#btn-back-0")
    pg.evaluate("S.parsedStories=[];S.implStep=0;S.generatedDocs=[];S.deployResults=[];refresh()")
    pg.click("#tab-text")
    # implement: extract -> review -> docs
    pg.fill("#story-input", "Route P1 incidents to on-call within 2 minutes. VPN is down in Pune.")
    pg.click("#btn-parse"); pg.wait_for_selector("text=Review User Stories", timeout=15000)
    ok("Extract stories via /snow/extract/universal", pg.locator(".story-row").count() == 4)
    pg.screenshot(path="shots/04_review.png")
    n_docs = pg.evaluate("DOC_TYPES.length")
    pg.click("#btn-gen-docs"); pg.wait_for_selector(f"text=All {n_docs} documents generated", timeout=60000)
    ok(f"All {n_docs} documents generated with their own templates", True)
    pg.screenshot(path="shots/05_docs.png")
    with pg.expect_download() as dl:
        pg.click("[data-dl='1'][data-fmt='docx']")
    ok("Download on-screen BRD as .docx", dl.value.suggested_filename.startswith("BRD"), dl.value.suggested_filename)
    pg.click("#btn-go-deploy")
    pg.click("[data-deploy='1']"); pg.click("#_yes")
    pg.wait_for_selector("text=✓ Done", timeout=15000)
    ok("Deploy one story after approval -> incident", "incident" in pg.inner_text("#mc"))
    pg.screenshot(path="shots/06_deploy.png")

    # chat: notes -> stories -> download from chat
    pg.click("[data-nav=chat]")
    pg.fill("#chat-input", "parse requirements: meeting notes - route P1 incidents to on-call within 2 minutes, VPN outage in Pune needs logging")
    pg.keyboard.press("Enter"); pg.wait_for_selector("text=/Extracted \\d+ user stories/", timeout=20000)
    with pg.expect_download() as dlc:
        pg.locator(".chat-action", has_text="User Stories (Word)").last.click()
    ok("Chat: extract from notes, then download User Stories (Word) from the chat", dlc.value.suggested_filename.endswith(".docx"), dlc.value.suggested_filename)

    # health scan
    pg.click("[data-nav=healthscan]"); pg.click("#btn-run-scan")
    pg.wait_for_selector("text=Category Scores", timeout=30000)
    txt = pg.inner_text("#mc")
    ok("Health scan shows live result (no hard-coded 74)", "142 CIs" not in txt and "OVERALL SCORE" in txt.upper() and "Scanned https://127.0.0.1:9443" in txt)
    pg.screenshot(path="shots/07_health.png", full_page=True)

    # build app
    pg.click("[data-nav=build]")
    pg.click("#b-gen-plan"); pg.wait_for_selector("text=Plan passed all checks", timeout=20000)
    pg.fill("#b-appname", "UI Router")
    pg.screenshot(path="shots/08_build_plan.png", full_page=True)
    pg.uncheck("#b-stop")
    pg.click("#b-build"); pg.click("#_cyes")
    pg.wait_for_selector("text=/Build (completed|failed)/", timeout=60000)
    st = pg.inner_text("#mc")
    acl_lines = [l for l in st.splitlines() if "ACL:" in l]
    honest = bool(acl_lines) and all(("❌" in l or "⚠️" in l) and "✅" not in l for l in acl_lines) \
        and any("not created" in l or "❌" in l for l in acl_lines)
    ok("Build app from UI finishes with an honest status (ACL not created is reported, not ✅)",
       "completed with issues" in st and honest, " | ".join(acl_lines)[:600])
    pg.wait_for_selector("text=created artifacts verified", timeout=10000)
    ok("UI shows update-set capture verification", True)
    pg.click("#b-changes"); time.sleep(1.5)
    pg.screenshot(path="shots/09_build_done.png", full_page=True)

    # operations: all tools + plain-English ask
    pg.click("[data-nav=ops]")
    pg.click("[data-opstab=tools]")
    pg.wait_for_selector("[data-tool='list_incidents']", timeout=20000)
    n_tools = pg.evaluate("fetch(API+'/snow/tools').then(r=>r.json()).then(d=>d.total)")
    ok(f"Operations lists the catalog's {n_tools} tools", f"({n_tools})" in pg.inner_text("#mc"), n_tools)
    pg.click("[data-tool='list_incidents']"); pg.wait_for_selector("#ops-run", timeout=10000)
    pg.click("#ops-run"); pg.wait_for_selector("text=/row\\(s\\)/", timeout=20000)
    ok("Run a tool from the form -> results table", "INC" in pg.inner_text("#mc"))
    pg.screenshot(path="shots/11_ops_tools.png", full_page=True)
    pg.fill("#ops-ask", "create a problem for printers failing"); pg.click("#ops-ask-go")
    pg.wait_for_selector("#_cyes", timeout=20000)
    ok("Plain-English write request asks for confirmation first", "create_problem" in pg.inner_text("#_modal"))
    pg.click("#_cyes"); pg.wait_for_selector("text=/created|success/i", timeout=20000)
    pg.screenshot(path="shots/12_ops_ask.png", full_page=True)
    pg.click("[data-opstab=approvals]"); pg.wait_for_selector("text=/Nothing waiting|Approve/", timeout=15000)
    ok("Approvals tab loads", True)
    pg.click("[data-opstab=knowledge]"); ok("Knowledge tab renders", pg.locator("#ops-kb-create").count() == 1)
    pg.click("[data-opstab=catalog]"); ok("Catalog tab renders", pg.locator("#ops-cat-search").count() == 1)

    # update sets
    pg.click("[data-nav=updatesets]")
    pg.wait_for_selector("text=/\\d+ change\\(s\\)/", timeout=15000)
    ok("Update Sets screen lists sets with current marker", "HAI_UI_Router" in pg.inner_text("#mc"))
    pg.screenshot(path="shots/10_updatesets.png", full_page=True)
    b.close()

real = [e for e in errors if "ERR_TUNNEL_CONNECTION_FAILED" not in e and "401" not in e]  # CDN blocked in sandbox; 401 = the deliberate wrong-password test
ok("No JavaScript errors in the page", not real, "; ".join(real[:5]))
print(f"==== {sum(1 for _, x in results if x)}/{len(results)} UI checks passed ====")
